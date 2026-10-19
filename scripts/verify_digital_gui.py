"""Explicit local Qt acceptance. Private images/screenshots stay in local-data."""

from __future__ import annotations

import argparse
import faulthandler
import hashlib
import io
import json
import re
import sys
import time
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image
from PySide6.QtCore import QEventLoop, QTimer
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QMessageBox, QPushButton, QTextEdit

from document_excel.infrastructure.paddle_ocr import _POOL, PaddleOCREngine
from document_excel.infrastructure.storage import Settings, SettingsRepository
from document_excel.presentation.main_window import MainWindow
from document_excel.presentation.theme import apply_theme


def wait_for(predicate, timeout=120):
    if predicate():
        return
    deadline = time.monotonic() + timeout
    loop = QEventLoop()
    timer = QTimer()
    timer.timeout.connect(lambda: loop.quit() if predicate() or time.monotonic() > deadline else None)
    timer.start(25)
    loop.exec()
    timer.stop()
    if not predicate():
        raise RuntimeError("Qt acceptance timeout")


def canonical(text):
    return re.sub(r"\W", "", str(text)).casefold()


def main():
    faulthandler.dump_traceback_later(60, repeat=True)
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--images", nargs=2, type=Path, required=True)
    parser.add_argument("--models", type=Path, required=True)
    parser.add_argument("--engine", choices=["rapidocr", "paddle"], default="rapidocr")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--inspect-only", action="store_true")
    args = parser.parse_args()
    expected = json.loads(sys.stdin.readline())
    root = Path(__file__).resolve().parents[1]
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    repo = SettingsRepository(output / "app-data")
    repo.save(
        Settings(
            ocr_engine=args.engine,
            rapid_model_directory=str(args.models.resolve()),
            paddle_model_directory=str(args.models.resolve()),
            timeout=120,
            default_template=str(root / "assets" / "lermontov.xlsx"),
        )
    )
    window_started = time.perf_counter()
    app = QApplication.instance() or QApplication([])
    app.setApplicationName("Local digital acceptance")
    apply_theme(app, repo.load())
    window = MainWindow(repo)
    window.resize(1366, 900)
    errors = []
    window.error = lambda *_: errors.append(True)
    window.show()
    report = {"engine": args.engine}
    calls = []
    native = PaddleOCREngine._read

    def read(self, *a, **kw):
        calls.append(kw.get("mode", a[4] if len(a) > 4 else "document"))
        return native(self, *a, **kw)

    PaddleOCREngine._read = read
    try:
        print("Qt: opening application", flush=True)
        wait_for(lambda: not window.background_workers)
        report["window_and_template_ready_seconds"] = time.perf_counter() - window_started
        window.case_actions.add_images(args.images)
        start = time.perf_counter()
        ticks = []
        heartbeat = QTimer(window)
        heartbeat.timeout.connect(lambda: ticks.append(time.perf_counter()))
        heartbeat.start(50)
        print("Qt: starting document recognition", flush=True)
        window.case_actions.start_ocr()
        wait_for(lambda: window.active_worker is None and not window.background_workers)
        report["ocr_seconds"] = time.perf_counter() - start
        heartbeat.stop()
        report["gui_heartbeat_ticks_during_ocr"] = len(ticks)
        print("Qt: recognition finished", flush=True)
        assert not errors
        case = window.current_case
        preview = window.case_actions.ocr_state.preview(case)
        assert preview and preview.data.full_name
        report["pending_fields_visible"] = bool(window.review_form.fields["full_name"].editor.text())
        report["review_page_visible_after_worker_finishes"] = (
            window.current_step == 2 and window.review_form.isVisible()
        )
        assert report["review_page_visible_after_worker_finishes"]
        report["pending_readonly"] = window.review_form.fields["surname"].editor.isReadOnly()
        report["pending_live_data_empty"] = case.data.full_name is None
        report["preview_exact_checks"] = {
            key: str(getattr(preview.data, key)) == value
            for key, value in expected.items()
            if key != "registration_address"
        }
        window.grab().save(str(output / "01-pending.png"))

        def accept_owner():
            dialog = QApplication.activeModalWidget()
            assert dialog
            for button in dialog.findChildren(QPushButton):
                if button.text() == "Подтверждаю: один человек":
                    button.click()
                    return
            raise RuntimeError("Identity confirmation button missing")

        calls_before_confirmation = len(calls)
        QTimer.singleShot(75, accept_owner)
        window.case_actions.review_identity()
        print("Qt: ownership confirmed", flush=True)
        assert case.data.full_name and case.identity_status == "user_confirmed"
        report["identity_confirmed_without_ocr"] = len(calls) == calls_before_confirmation
        before = case.data.model_dump()
        count = len(calls)
        for index in (1, 0, 1, 0):
            window.image_list.setCurrentIndex(index)
            QTest.qWait(60)
        wait_for(lambda: not window.background_workers)
        assert len(calls) == count and case.data.model_dump() == before
        report["switching_preserves_data_without_ocr"] = True
        window.grab().save(str(output / "02-fields.png"))
        report["exact_checks"] = {
            key: str(getattr(case.data, key)) == value
            for key, value in expected.items()
            if key != "registration_address"
        }
        report["address_candidates"] = [
            {
                "method": c.extraction_method,
                "control_match": canonical(c.normalized_value)
                == canonical(expected.get("registration_address", "")),
            }
            for c in case.evidence["registration_address"].candidates
        ]

        def capture_address():
            dialog = QApplication.activeModalWidget()
            assert dialog
            dialog.grab().save(str(output / "03-address.png"))
            dialog.reject()

        QTimer.singleShot(150, capture_address)
        window.show_field_details("registration_address")
        print("Qt: address alternatives checked", flush=True)

        def capture_safe():
            dialog = QApplication.activeModalWidget()
            assert dialog
            editor = dialog.findChild(QTextEdit)
            body = editor.toPlainText()
            report["masked_diagnostics_no_values"] = not any(
                value in body
                for key, value in expected.items()
                if key in {"personal_identifier", "full_name", "document_number"}
            )
            report["diagnostic_stage_metrics_present"] = "text_detection_seconds" in body
            dialog.grab().save(str(output / "04-diagnostics-masked.png"))
            dialog.accept()

        QTimer.singleShot(150, capture_safe)
        window.show_raw_ocr()
        print("Qt: masked diagnostics checked", flush=True)
        if not args.inspect_only:
            before_field = case.data.model_dump()
            count = len(calls)
            candidate = case.evidence["birth_date"].candidates[0]
            window.case_actions.repeat_field("birth_date", candidate)
            print("Qt: single-field worker started", flush=True)
            wait_for(lambda: window.active_worker is None and not window.background_workers)
            print("Qt: single-field worker finished", flush=True)
            report["single_field_retry_preserves_data"] = case.data.model_dump() == before_field
            report["single_field_retry_requests"] = len(calls) - count
            assert report["single_field_retry_preserves_data"] and 0 < len(calls) - count <= 3
        report["ok"] = all(report["exact_checks"].values()) and all(report["preview_exact_checks"].values())
        report["full_excel_scenario"] = False
        if not args.inspect_only:
            # An occluded address remains a review task. Exercise the actual
            # address editor with the independently transcribed local control;
            # this is recorded as manual correction, never OCR accuracy.
            def correct_address():
                dialog = QApplication.activeModalWidget()
                from document_excel.domain.digital_address import digital_address_from_text

                parsed = digital_address_from_text(expected["registration_address"])
                for name, editor in dialog.entries.items():
                    editor.setText(getattr(parsed, name))
                dialog.raw.setText(expected["registration_address"])
                report["control_address_parsed_fields"] = {
                    name: bool(editor.text()) for name, editor in dialog.entries.items()
                }
                dialog.accept()
                if dialog.isVisible():
                    report["address_form_validation"] = dialog.error.text()
                    dialog.grab().save(str(output / "address-validation.png"))
                    dialog.reject()

            QTimer.singleShot(75, correct_address)
            window.case_actions.address_dialog(True)
            print("Qt: address reviewed", flush=True)
            report["address_manually_reviewed"] = True

            def confirm_values():
                dialog = QApplication.activeModalWidget()
                assert dialog
                if isinstance(dialog, QMessageBox):
                    dialog.button(QMessageBox.StandardButton.Yes).click()
                    return
                for button in dialog.findChildren(QPushButton):
                    if button.text() in {"Подтвердить", "Да"}:
                        button.click()
                        return
                dialog.accept()

            QTimer.singleShot(75, confirm_values)
            window.case_actions.confirm_data()
            print("Qt: values confirmed", flush=True)
            assert case.confirmed and not errors
            source = root / "assets/lermontov.xlsx"
            source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
            plan = window.workbooks.plan(source, case, window.profile)
            target = window.workbooks.write(plan, output / "digital-contract.xlsx", case, window.profile)
            wb = load_workbook(target)
            original_wb = load_workbook(source)
            checks = {}
            for change in plan.changes:
                actual = wb[change.sheet][change.cell].value
                wanted = change.new
                if hasattr(actual, "date") and hasattr(wanted, "year"):
                    actual = actual.date()
                checks[change.sheet + "!" + change.cell] = actual == wanted
            checks["source_unchanged"] = hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
            changes = {(c.sheet, c.cell) for c in plan.changes}
            checks["formulas_preserved"] = all(
                wb[sheet.title][cell.coordinate].value == cell.value
                for sheet in original_wb
                for row in sheet
                for cell in row
                if cell.data_type == "f" and (sheet.title, cell.coordinate) not in changes
            )
            checks["styles_preserved"] = all(
                wb[sheet.title][cell.coordinate]._style == cell._style
                for sheet in original_wb
                for row in sheet
                for cell in row
                if cell.has_style
            )
            photos = wb[window.profile.photo_sheet]._images
            # Preserve the template's unrelated logo outside the photo ranges.
            # Identify our two photographs by their exact inserted PNG bytes.
            inserted = {hashlib.sha256(p.data).hexdigest(): p for p in plan.photos.insertions}
            matched = [
                (p, inserted[digest])
                for p in photos
                if (digest := hashlib.sha256(p._data()).hexdigest()) in inserted
            ]
            checks["two_current_client_photos"] = len(matched) == 2

            def anchor_key(anchor):
                return tuple(
                    getattr(marker, key)
                    for marker in (anchor._from, anchor.to)
                    for key in ("col", "row", "colOff", "rowOff")
                )

            checks["photo_anchors_match_plan"] = (
                all(anchor_key(p.anchor) == anchor_key(spec.anchor) for p, spec in matched)
                and len(matched) == 2
            )
            checks["photos_inside_target_ranges"] = (
                all(
                    (0 if spec.role.value == "FRONT" else 5) <= p.anchor._from.col
                    and p.anchor.to.col < (5 if spec.role.value == "FRONT" else 11)
                    and p.anchor._from.row >= 0
                    and p.anchor.to.row < 24
                    for p, spec in matched
                )
                and len(matched) == 2
            )
            checks["unrelated_template_photos_preserved"] = (
                len(photos) == len(original_wb[window.profile.photo_sheet]._images) + 2
            )
            from document_excel.infrastructure.excel_images import SheetGeometry

            geometry = SheetGeometry(wb[window.profile.photo_sheet])
            proportions = []
            for photo, spec in matched:
                start, end = photo.anchor._from, photo.anchor.to
                x0, y0 = geometry.position(start.col, start.row, start.colOff, start.rowOff)
                x1, y1 = geometry.position(end.col, end.row, end.colOff, end.rowOff)
                with Image.open(io.BytesIO(spec.data)) as pixels:
                    proportions.append(abs((x1 - x0) / (y1 - y0) - pixels.width / pixels.height) < 1e-5)
            checks["photo_aspect_preserved"] = len(proportions) == 2 and all(proportions)
            checks["pin_is_text"] = wb["Дог."]["C22"].data_type == "s"
            checks["number_is_text"] = wb["Дог."]["G19"].data_type == "s"
            checks["fio_matches_ocr"] = wb["Дог."]["A18"].value == case.data.full_name
            checks["address_matches_review"] = canonical(case.data.registration_address.raw) == canonical(
                expected["registration_address"]
            )
            report["excel_checks"] = checks
            report["full_excel_scenario"] = all(checks.values())
            report["ok"] = report["ok"] and report["full_excel_scenario"]
            window.grab().save(str(output / "05-export.png"))
            print("Qt: Excel reopened and checked", flush=True)
    finally:
        PaddleOCREngine._read = native
        wait_for(lambda: not window.background_workers)
        window.close()
        QTest.qWait(100)
        _POOL.close()
        (output / "report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report))
    faulthandler.cancel_dump_traceback_later()
    return 0 if report.get("ok") else 1


if __name__ == "__main__":
    raise SystemExit(main())
