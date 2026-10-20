"""Real offline CPU OCR through the customer GUI, using only generated data."""

from __future__ import annotations

import argparse
import hashlib
import json
import time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image
from PySide6.QtCore import QEventLoop, QTimer
from PySide6.QtWidgets import QApplication, QDialog, QFileDialog

from document_excel.domain.models import Address, FieldEvidence, FieldStatus, RegistrationStatus
from document_excel.infrastructure.digital_verification import synthetic_digital_images
from document_excel.infrastructure.excel import create_demo_template, file_hash
from document_excel.infrastructure.storage import Settings, SettingsRepository
from document_excel.presentation.components import ConfirmDialog
from document_excel.presentation.dialogs import ExportDialog
from document_excel.presentation.main_window import MainWindow
from document_excel.presentation.theme import apply_theme


def wait_for(predicate, timeout=180):
    deadline = time.monotonic() + timeout
    loop = QEventLoop()
    timer = QTimer()
    timer.timeout.connect(lambda: loop.quit() if predicate() or time.monotonic() > deadline else None)
    timer.start(20)
    if not predicate():
        loop.exec()
    timer.stop()
    if not predicate():
        raise RuntimeError("GUI verification timed out")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--models", type=Path, required=True)
    parser.add_argument("--engine", choices=["rapidocr", "paddle"], default="rapidocr")
    args = parser.parse_args()
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    source = output / "synthetic-filled-source.xlsx"
    create_demo_template(source)
    book = load_workbook(source)
    for address, value in {
        "A18": "Прежний Тестовый Клиент",
        "C22": "00987654321098",
        "E19": "ID",
        "G19": "0000002",
        "J19": "МКК",
        "K19": "999999",
        "M19": "01.01.2010",
        "A23": "01.01.1980",
        "C20": "Прежний регион",
        "J20": "Прежняя улица",
        "D7": "+996 555 000001",
    }.items():
        book["Дог."][address] = value
    for address, value in {"B20": "Синтетический перфоратор", "I20": 2, "N20": 350, "P20": 5000}.items():
        book["Пр №1"][address] = value
    book["Дог."]["A21"] = "=C20"
    book["Дог."]["I21"] = "=J20"
    photo_sheet = book.create_sheet("Пасп Доп Данные")
    photo_sheet["M30"] = "Логотип"
    for index, anchor in enumerate(("A1", "F1", "M1")):
        path = output / f"old-synthetic-{index}.png"
        Image.new("RGB", (120, 80), (100, index * 80, 20)).save(path)
        photo_sheet.add_image(ExcelImage(path), anchor)
    book.save(source)
    book.close()
    original_hash = file_hash(source)
    paths, expected = synthetic_digital_images(output / "generated-passport")
    repo = SettingsRepository(output / "app-data")
    repo.save(
        Settings(
            ocr_engine=args.engine,
            rapid_model_directory=str(args.models.resolve()),
            paddle_model_directory=str(args.models.resolve()),
            timeout=120,
            output_directory=str(output / "results"),
        )
    )
    app = QApplication.instance() or QApplication([])
    apply_theme(app, repo.load())
    window = MainWindow(repo)
    errors = []
    window.error = lambda message: errors.append(str(message))
    window.show()
    window.grab().save(str(output / "01-start.png"))
    QFileDialog.getOpenFileName = lambda *args, **kwargs: (str(source), "")
    window.contract_button.click()
    wait_for(lambda: not window.background_workers)
    window.case_actions.add_images(paths)
    wait_for(lambda: not window.background_workers)
    for width, height in ((1366, 900), (1000, 700), (800, 600)):
        window.resize(width, height)
        app.processEvents()
        window.grab().save(str(output / f"02-upload-{width}.png"))
    window.resize(1366, 900)
    assert window.can_recognize()
    ticks = []
    timer = QTimer()
    timer.timeout.connect(lambda: ticks.append(time.monotonic()))
    timer.start(20)
    started = time.monotonic()
    window.primary_button.click()
    wait_for(lambda: not window.background_workers and window.active_worker is None)
    elapsed = time.monotonic() - started
    timer.stop()
    case = window.current_case
    if case.pending_identity:
        original_exec = QDialog.exec
        QDialog.exec = lambda self: QDialog.DialogCode.Accepted
        try:
            window.case_actions.review_identity()
        finally:
            QDialog.exec = original_exec
    checks = {name: str(getattr(case.data, name)) == value for name, value in expected.items()}
    if not all(checks.values()):
        failure = {
            "ok": False,
            "engine": args.engine,
            "synthetic_only": True,
            "ocr_checks": checks,
            "ocr_seconds": round(elapsed, 3),
            "errors": case.errors,
            "metrics": case.processing_metrics,
            "status": case.status,
        }
        (output / "report.json").write_text(
            json.dumps(failure, ensure_ascii=False, indent=2), encoding="utf8"
        )
        window.close()
    assert all(checks.values()), checks
    # Explicit manual registration is a normal customer action, separate from
    # the accuracy assertions above and independent of address OCR confidence.
    case.data.registration_address = Address(region="Тестовая область", street="улица Примерная, дом 7")
    case.registration_status = RegistrationStatus.MANUALLY_ENTERED
    case.evidence["registration_address"] = FieldEvidence(
        source="ручной ввод", status=FieldStatus.MANUALLY_CORRECTED
    )
    case.invalidate()
    window.refresh_form()
    window.go_step(2)
    wait_for(lambda: window.viewer_panel.isVisible() and window.review_form.isVisible())
    app.processEvents()
    window.grab().save(str(output / "03-review.png"))
    ConfirmDialog.ask = staticmethod(lambda *args: True)

    def accept_export(dialog):
        dialog.confirm.setChecked(True)
        dialog.image_confirm.setChecked(True)
        dialog.clear_confirm.setChecked(True)
        dialog.preserved_confirm.setChecked(True)
        dialog.clear_contacts.setChecked(True)
        dialog.grab().save(str(output / "04-preview.png"))
        dialog.accept()
        return QDialog.DialogCode.Accepted

    ExportDialog.exec = accept_export
    QFileDialog.getSaveFileName = lambda *args, **kwargs: (args[2], "")
    window.primary_button.click()
    wait_for(lambda: case.output_path is not None or bool(errors))
    wait_for(lambda: not window.background_workers)
    assert not errors, errors
    wait_for(lambda: window.current_step == 4)
    app.processEvents()
    window.grab().save(str(output / "05-success.png"))
    result = load_workbook(case.output_path)
    preservation = {
        "tools": result["Пр №1"]["B20"].value == "Синтетический перфоратор",
        "price": result["Пр №1"]["N20"].value == 350,
        "quantity": result["Пр №1"]["I20"].value == 2,
        "deposit": result["Пр №1"]["P20"].value == 5000,
        "formula": result["Дог."]["A21"].value == "=C20",
        "name": result["Дог."]["A18"].value == "Синтетиков Пример Тестович",
        "pin": result["Дог."]["C22"].value == expected["personal_identifier"],
        "document": result["Дог."]["G19"].value == "0000001",
        "photos_plus_logo": len(result["Пасп Доп Данные"]._images) == 3,
        "old_phone_cleared": result["Дог."]["D7"].value is None,
        "original_unchanged": file_hash(source) == original_hash,
    }
    result.close()
    report = {
        "engine": args.engine,
        "synthetic_only": True,
        "ocr_checks": checks,
        "preservation": preservation,
        "ocr_seconds": round(elapsed, 3),
        "gui_timer_ticks": len(ticks),
        "output": str(case.output_path),
        "output_sha256": hashlib.sha256(case.output_path.read_bytes()).hexdigest(),
        "manual_registration": True,
        "ok": all(checks.values()) and all(preservation.values()) and len(ticks) > 5,
    }
    (output / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf8")
    window.close()
    app.processEvents()
    print(json.dumps(report, ensure_ascii=True))
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
