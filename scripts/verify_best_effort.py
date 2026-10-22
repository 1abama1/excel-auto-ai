"""Exercise the actual three-stage Qt workflow; use synthetic documents only."""

from __future__ import annotations

import argparse
import json
import os
import time
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication, QDialog

from document_excel.domain.models import DocumentSide
from document_excel.infrastructure.digital_verification import synthetic_digital_images
from document_excel.infrastructure.excel import create_demo_template, file_hash
from document_excel.infrastructure.storage import SettingsRepository
from document_excel.presentation.dialogs import CopySaveDialog
from document_excel.presentation.main_window import MainWindow
from document_excel.presentation.theme import apply_theme


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    parser.add_argument("--real-ocr", action="store_true")
    args = parser.parse_args()
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    app = QApplication([])
    window = MainWindow(SettingsRepository(output / "settings"))
    apply_theme(app, window.settings)
    window.resize(1366, 820)
    window.show()
    errors = []
    window.error = lambda text: errors.append(str(text))

    def wait():
        start = time.monotonic()
        while window.background_workers:
            app.processEvents()
            time.sleep(0.005)
            if time.monotonic() - start > 150:
                raise RuntimeError("Background operation timed out")
        app.processEvents()
        if errors:
            raise RuntimeError("Workflow failed: " + "; ".join(errors))

    wait()
    source = output / "synthetic-input.xlsx"
    create_demo_template(source)
    book = load_workbook(source)
    book.create_sheet("Пасп Доп Данные")["M30"] = "Synthetic logo area"
    book["Пр №1"]["B20"] = "Synthetic tool"
    book["Пр №1"]["N20"] = 350
    book["Пр №1"]["P20"] = 5000
    book.save(source)
    book.close()
    original_hash = file_hash(source)
    window.excel_actions.load_template(source)
    wait()
    for width, height in [(1366, 768), (900, 700), (800, 560)]:
        window.resize(width, height)
        app.processEvents()
        window.grab().save(str(output / f"upload-{width}.png"))
    window.resize(1366, 820)
    ocr_elapsed = None
    if args.real_ocr:
        paths, _ = synthetic_digital_images(output / "synthetic-photos")
        window.case_actions.add_images(paths[:1])
        wait()
        window.current_case.images[0].role = DocumentSide.FRONT
        window.refresh_form()
        assert window.can_recognize()
        started = time.monotonic()
        window.primary_button.click()
        wait()
        ocr_elapsed = round(time.monotonic() - started, 3)
        assert window.current_step == 2
        assert window.current_case.data.full_name
    else:
        window.skip_ocr_button.click()
    for width, height in [(1366, 768), (900, 700), (800, 560)]:
        window.resize(width, height)
        app.processEvents()
        window.grab().save(str(output / f"review-{width}.png"))
    window.resize(1366, 820)
    old_exec = CopySaveDialog.exec
    dialogs = []

    def accept(dialog):
        dialogs.append(True)
        assert dialog.in_place and dialog.target.parent == source.parent
        dialog.show()
        app.processEvents()
        dialog.grab().save(str(output / "save-summary.png"))
        assert dialog.save_button.isEnabled()
        dialog.accept()
        return QDialog.DialogCode.Accepted

    CopySaveDialog.exec = accept
    try:
        window.primary_button.click()
        wait()
        # The worker dispatch is deferred until the plan worker has finished.
        for _ in range(30):
            app.processEvents()
            if window.active_worker:
                wait()
            if window.current_case.output_path:
                break
            time.sleep(0.01)
    finally:
        CopySaveDialog.exec = old_exec
    result = window.current_case.output_path
    assert result and result.exists() and len(dialogs) == 1
    app.processEvents()
    assert window.current_step == 4
    window.grab().save(str(output / "success.png"))
    check = load_workbook(result)
    checks = {
        "selected_file_updated": file_hash(result) != original_hash
        if args.real_ocr
        else file_hash(result) == original_hash,
        "same_folder_renamed": result.parent == source.parent and not source.exists(),
        "selected_path_updated": window.current_case.template_path == result,
        "no_extra_xlsx": list(output.glob("*.xlsx")) == [result],
        "commercial_data_preserved": check["Пр №1"]["N20"].value == 350
        and check["Пр №1"]["P20"].value == 5000,
        "formula_preserved": check["Акт расч"]["L21"].value == "=G21-F21",
        "single_final_dialog": len(dialogs) == 1,
        "three_visible_steps": sum(b.isVisible() for b in window.step_indicator.buttons) == 3,
        "no_sidebar": window.sidebar.isHidden(),
        "one_photo_only": len(window.current_case.images) == 1
        if args.real_ocr
        else not window.current_case.images,
        "success_page": window.current_step == 4,
    }
    if args.real_ocr:
        checks["name_written"] = check["Дог."]["A18"].value == window.current_case.formatted_name
        checks["missing_registration_empty"] = check["Дог."]["J20"].value is None
        checks["front_photo_inserted"] = len(check["Пасп Доп Данные"]._images) == 1
    check.close()
    report = {
        "ok": all(checks.values()),
        "output": str(result),
        "checks": checks,
        "ocr_seconds": ocr_elapsed,
        "updated_fields": len(window.current_case.updated_fields),
        "missing_fields": len(window.current_case.missing_fields),
    }
    (output / "report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report))
    window.close()
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
