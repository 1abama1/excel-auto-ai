from __future__ import annotations

import errno
import hashlib
import os
import re
import shutil
import stat
import tempfile
import zipfile
from copy import deepcopy
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any, Literal

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.worksheet.print_settings import PrintArea
from pydantic import BaseModel, Field, field_validator, model_validator

from document_excel.domain.errors import AppError, OutputExistsError
from document_excel.domain.models import (
    FIELD_LABELS,
    ClientCase,
    ContractData,
    DocumentSide,
    FieldStatus,
    MissingValuePolicy,
    RegistrationStatus,
    WorkflowMode,
    data_value,
    validate_field,
)
from document_excel.infrastructure.contract_filename import ContractFilenameGenerator
from document_excel.infrastructure.excel_images import ExcelImageEmbedder, PhotoExportPlan
from document_excel.infrastructure.storage import atomic_json
from document_excel.infrastructure.xlsx_package import PackageEditor, remap_formula

BASE_FIELDS = set(FIELD_LABELS) | {"document_type"}
ADDRESS_FIELDS = {
    f"{prefix}.{part}"
    for prefix in ("registration_address", "living_address")
    for part in ("region", "street", "full")
}
EXPORT_FIELDS = (
    BASE_FIELDS
    | ADDRESS_FIELDS
    | {
        "phones",
        "tool_full_name",
        "purchase_price",
        "deposit_amount",
        "actual_return_date",
        "actual_return_time",
    }
)


class CellMapping(BaseModel):
    field: str
    sheet: str
    cell: str
    value_type: Literal["text", "date", "number"] = "text"
    allow_formula_replace: bool = False
    skip_empty: bool = True

    @field_validator("field")
    @classmethod
    def known_field(cls, value: str) -> str:
        if value not in EXPORT_FIELDS:
            raise ValueError("Неизвестное поле данных")
        return value

    @field_validator("cell")
    @classmethod
    def valid_cell(cls, value: str) -> str:
        value = value.upper().strip()
        if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", value):
            raise ValueError("Некорректный адрес ячейки")
        row, col = coordinate_to_tuple(value)
        if row > 1048576 or col > 16384:
            raise ValueError("Адрес вне диапазона Excel")
        return value


class TemplateProfile(BaseModel):
    id: str = "custom"
    name: str = "Собственный профиль"
    version: str = "1"
    structure_digest: str = ""
    required_sheets: list[str] = Field(default_factory=lambda: ["Дог."])
    optional_sheets: list[str] = Field(default_factory=list)
    field_mappings: list[CellMapping] = Field(default_factory=list)
    required_fields: list[str] = Field(
        default_factory=lambda: [
            "full_name",
            "personal_identifier",
            "passport_number",
            "registration_address",
        ]
    )
    pin_length: int = Field(default=14, ge=6, le=30)
    phone_mode: Literal["combined", "separate"] = "combined"
    allow_missing_registration: bool = False
    identifier_label: str = "ПИН / ИНН"
    tool_sheet: str | None = "Пр №1"
    tool_first_row: int = Field(default=20, ge=1, le=1000000)
    tools_per_page: int = Field(default=5, ge=1, le=100)
    tool_row_step: int = Field(default=1, ge=1, le=100)
    tool_inventory_column: str | None = "K"
    tool_name_formula_replace: bool = False
    tool_linked_sheets: list[str] = Field(default_factory=list)
    tool_inventory_sheet: str | None = None
    tool_inventory_first_row: int = Field(default=21, ge=1, le=1000000)
    tool_inventory_target_column: str = "L"
    enable_return_formula: bool = False
    acceptance_sheet: str | None = None
    photo_sheet: str | None = None
    photos_required: bool = False
    allow_photo_exception: bool = False

    @model_validator(mode="after")
    def consistent(self) -> TemplateProfile:
        for column in (self.tool_inventory_column, self.tool_inventory_target_column):
            if column is not None:
                CellMapping.valid_cell(column + "1")
        addresses = [(m.sheet, m.cell) for m in self.field_mappings]
        if len(set(addresses)) != len(addresses):
            raise ValueError("Два поля назначены одной ячейке")
        if any(name not in BASE_FIELDS for name in self.required_fields):
            raise ValueError("Неизвестное обязательное поле")
        fields = {m.field for m in self.field_mappings}
        if self.phone_mode == "combined" and fields & {"whatsapp_phone", "additional_phone"}:
            raise ValueError("Для объединённых телефонов используйте единственное поле phones")
        if self.phone_mode == "separate" and "phones" in fields:
            raise ValueError("Для раздельных телефонов используйте whatsapp_phone и additional_phone")
        return self


def legacy_profile(phone_mode: Literal["combined", "separate"] = "combined") -> TemplateProfile:
    pairs = {
        "tool_full_name": "A12",
        "full_name": "A18",
        "passport_type": "E19",
        "passport_number": "G19",
        "passport_issued_by": "J19",
        "passport_department_code": "K19",
        "passport_issued_date": "M19",
        "registration_address.region": "C20",
        "registration_address.street": "J20",
        "living_address.region": "A21",
        "living_address.street": "I21",
        "personal_identifier": "C22",
        "birth_date": "A23",
        "object_address": "K22",
    }
    pairs.update(
        {"phones": "D7"} if phone_mode == "combined" else {"whatsapp_phone": "D7", "additional_phone": "G7"}
    )
    mappings = [CellMapping(field=name, sheet="Дог.", cell=cell) for name, cell in pairs.items()]
    for name, cell in [
        ("actual_return_date", "G21"),
        ("actual_return_time", "G20"),
        ("purchase_price", "N21"),
    ]:
        mappings.append(
            CellMapping(
                field=name,
                sheet="Акт расч",
                cell=cell,
                value_type="date"
                if name == "actual_return_date"
                else "number"
                if name == "purchase_price"
                else "text",
            )
        )
    return TemplateProfile(
        id="legacy-" + phone_mode,
        name="Карта из задания · " + phone_mode,
        phone_mode=phone_mode,
        field_mappings=mappings,
        optional_sheets=[
            "Акт расч",
            "Акт прием",
            "Акт приема",
            "Акт прием-передач",
            "Акт приема-передачи",
            "Акт прием-передачи",
        ],
    )


def lermontov_profile() -> TemplateProfile:
    """Verified against the supplied assets/lermontov.xlsx, not its filename."""
    profile = legacy_profile()
    profile.id = "lermontov-combined"
    profile.name = "Договор Лермонтова · проверенная карта"
    profile.version = "3"
    profile.photo_sheet = "Пасп Доп Данные"
    profile.photos_required = True
    profile.allow_photo_exception = True
    profile.tool_row_step = 2
    profile.tool_inventory_column = None
    profile.tool_name_formula_replace = True
    profile.tool_linked_sheets = ["Акт прием", "Акт расч"]
    profile.tool_inventory_sheet = "Акт прием"
    profile.acceptance_sheet = "Акт прием"
    profile.field_mappings = [m for m in profile.field_mappings if m.sheet == "Дог."]
    for mapping in profile.field_mappings:
        if mapping.field.startswith("living_address."):
            mapping.allow_formula_replace = True
        if mapping.field in {"birth_date", "passport_issued_date"}:
            mapping.value_type = "date"
    profile.field_mappings.append(
        CellMapping(
            field="purchase_price",
            sheet="Акт прием",
            cell="N21",
            value_type="number",
            allow_formula_replace=True,
        )
    )
    for row in range(20, 30, 2):
        for name, address, kind in [
            ("actual_return_time", f"G{row}", "text"),
            ("actual_return_date", f"G{row + 1}", "date"),
        ]:
            profile.field_mappings.append(
                CellMapping(
                    field=name, sheet="Пр №1", cell=address, value_type="date" if kind == "date" else "text"
                )
            )
    return profile


class TemplateRegistry:
    def __init__(self, root: Path) -> None:
        self.root = root

    def save(self, profile: TemplateProfile) -> Path:
        # Filename never derives from an imported id.
        path = self.root / f"profile-{hashlib.sha256(profile.id.encode()).hexdigest()[:16]}.json"
        atomic_json(path, profile.model_dump())
        return path

    def load(self, path: Path) -> TemplateProfile:
        try:
            if path.stat().st_size > 1_000_000:
                raise AppError("Профиль слишком большой.")
            return TemplateProfile.model_validate_json(path.read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            raise AppError("Профиль mapping некорректен. Проверьте поля, ячейки и телефонный режим.") from exc

    def matching(self, book) -> TemplateProfile | None:
        for path in sorted(self.root.glob("profile-*.json")):
            try:
                profile = self.load(path)
                if profile.structure_digest and profile.structure_digest == workbook_structure(book, profile):
                    return profile
            except AppError:
                continue
        return None


@dataclass
class RecognitionWarnings:
    found_fields: list[str] = field(default_factory=list)
    missing_fields: list[str] = field(default_factory=list)
    review_fields: list[str] = field(default_factory=list)
    messages: list[str] = field(default_factory=list)


@dataclass
class ExportSafetyErrors:
    errors: list[str] = field(default_factory=list)


# Names intentionally alias the two public validation results.
ValidationWarnings = RecognitionWarnings
BlockingExportErrors = ExportSafetyErrors


def usable_value(case: ClientCase, name: str, pin_length: int = 14) -> bool:
    value = case.formatted_name if name == "full_name" else data_value(case.data, name)
    evidence = case.evidence.get(name)
    if not value or value.strip().casefold() in {"none", "unknown", "n/a", "не найдено", "не указано", "0"}:
        return False
    if evidence and evidence.status in {FieldStatus.INVALID, FieldStatus.CONFLICT}:
        return False
    if validate_field(name, value, pin_length):
        return False
    if name == "full_name" and not re.fullmatch(r"[^\W\d_]+(?:[ '’-][^\W\d_]+)*", value):
        return False
    return True


def recognition_warnings(case: ClientCase, profile: TemplateProfile) -> RecognitionWarnings:
    result = RecognitionWarnings()
    names = list(
        dict.fromkeys(
            m.field.split(".")[0]
            for m in profile.field_mappings
            if m.field in EXISTING_FIELDS
            and m.field != "object_address"
            and m.field.split(".")[0] in FIELD_LABELS
        )
    )
    for name in names:
        if not usable_value(case, name, profile.pin_length):
            result.missing_fields.append(name)
        else:
            result.found_fields.append(name)
        evidence = case.evidence.get(name)
        if evidence and evidence.status in {FieldStatus.INVALID, FieldStatus.CONFLICT, FieldStatus.REVIEW}:
            result.review_fields.append(name)
    if result.missing_fields:
        result.messages.append(
            "Не удалось определить: " + ", ".join(FIELD_LABELS[n] for n in result.missing_fields)
        )
    if result.review_fields:
        result.messages.append("Некоторые значения требуют проверки по фотографии.")
    if any(e.source_method == "MANUAL_INPUT" for e in case.evidence.values()):
        result.messages.append("Некоторые данные изменены вручную.")
    sides = {i.role for i in case.images if i.id not in case.side_proposals}
    if not {DocumentSide.FRONT, DocumentSide.BACK} <= sides:
        result.messages.append(
            "Одна или обе фотографии отсутствуют; изображения могут остаться без обновления."
        )
    if (
        case.side_proposals
        or any(i.role.value == "UNKNOWN" for i in case.images)
        or any(
            sum(i.role == side for i in case.images) > 1 for side in (DocumentSide.FRONT, DocumentSide.BACK)
        )
    ):
        result.messages.append("Уточните стороны, если хотите вставить эти фотографии в договор.")
    if case.identity_status == "unverified":
        result.messages.append("Проверьте принадлежность фотографий одному клиенту в итоговой сводке.")
    if case.data.document_type == "unknown" and case.images:
        result.messages.append("Тип документа не определён; доступные сведения можно сохранить.")
    return result


def export_safety_errors(case: ClientCase, profile: TemplateProfile) -> ExportSafetyErrors:
    errors = []
    if not case.template_path:
        errors.append("Сначала выберите готовый Excel-договор.")
    if case.pending_identity or case.identity_status == "review":
        errors.append("Проверьте принадлежность фотографий одному человеку перед сохранением.")
    if not profile.field_mappings:
        errors.append("Один раз настройте расположение полей этого договора в дополнительных параметрах.")
    return ExportSafetyErrors(errors)


def readiness(case: ClientCase, profile: TemplateProfile, require_confirmation: bool = True) -> list[str]:
    if case.workflow_mode == WorkflowMode.BEST_EFFORT_UPDATE:
        return export_safety_errors(case, profile).errors
    errors = []
    if (
        case.workflow_mode == WorkflowMode.EXISTING_CONTRACT
        and case.data.full_name
        and (
            not re.search(r"[А-Яа-яЁёҮүӨөҢң]", case.data.full_name)
            or not re.fullmatch(r"[^\W\d_]+(?:[ '\u2019-][^\W\d_]+)*", " ".join(case.data.full_name.split()))
        )
    ):
        errors.append("Для договора укажите подтверждённое ФИО на кириллице.")
    if case.side_proposals:
        errors.append("Проверьте определённые стороны документа перед сохранением фотографий.")
    for name in profile.required_fields:
        value = data_value(case.data, name)
        if (
            name == "registration_address"
            and not value
            and profile.allow_missing_registration
            and case.registration_status == RegistrationStatus.NOT_PRESENT
        ):
            continue
        error = validate_field(name, value, profile.pin_length)
        if error:
            errors.append(f"{FIELD_LABELS.get(name, name)}: {error}")
    for name in FIELD_LABELS:
        value = data_value(case.data, name)
        if value:
            error = validate_field(name, value, profile.pin_length)
            if error:
                errors.append(f"{FIELD_LABELS[name]}: {error}")
    data = case.data
    if case.contract.actual_return_time and not re.fullmatch(
        r"(?:[01][0-9]|2[0-3]):[0-5][0-9]", case.contract.actual_return_time
    ):
        errors.append("Время возврата должно иметь формат ЧЧ:ММ.")
    if data.birth_date and data.passport_issued_date and data.passport_issued_date < data.birth_date:
        errors.append("Дата выдачи раньше даты рождения.")
    if (
        data.expiration_date
        and data.passport_issued_date
        and data.expiration_date < data.passport_issued_date
    ):
        errors.append("Срок действия заканчивается раньше даты выдачи.")
    if "registration_address" in profile.required_fields:
        allowed_absence = (
            profile.allow_missing_registration and case.registration_status == RegistrationStatus.NOT_PRESENT
        )
        if not allowed_absence and case.registration_status not in {
            RegistrationStatus.USER_CONFIRMED,
            RegistrationStatus.MANUALLY_ENTERED,
        }:
            errors.append("Подтвердите адрес регистрации или заполните его вручную.")
    for name, evidence in case.evidence.items():
        if evidence.status in {FieldStatus.CONFLICT, FieldStatus.INVALID}:
            errors.append(f"{FIELD_LABELS.get(name, name)}: устраните конфликт или ошибку.")
    if case.identity_status == "review":
        errors.append(
            "Подтвердите принадлежность всех фотографий одному человеку или разнесите документы по клиентам."
        )
    if data.document_type == "unknown" and case.images:
        errors.append("Выберите тип документа: автоматическое определение неоднозначно.")
    if require_confirmation and not case.confirmed:
        errors.append("Нажмите «Подтвердить данные» после сверки с документом.")
    return list(dict.fromkeys(errors))


class ExcelDataAssembler:
    def assemble(self, case: ClientCase, pin_length: int = 14) -> dict[str, Any]:
        result: dict[str, Any] = {name: getattr(case.data, name) for name in BASE_FIELDS}
        if case.workflow_mode.updates_existing:
            result["full_name"] = case.formatted_name
        for prefix in ("registration_address", "living_address"):
            address = getattr(case.data, prefix)
            result[prefix] = address.display()
            result[f"{prefix}.region"] = address.region
            if case.data.document_type == "kg_digital_id" and not address.region:
                result[f"{prefix}.region"] = address.city
            result[f"{prefix}.street"] = address.legacy_street()
            result[f"{prefix}.full"] = address.display()
        result["phones"] = "; ".join(
            dict.fromkeys(p for p in [case.data.whatsapp_phone, case.data.additional_phone] if p)
        )
        result["tool_full_name"] = ", ".join(tool.name for tool in case.contract.tools)
        result.update(case.contract.model_dump(exclude={"tools"}))
        if case.workflow_mode == WorkflowMode.BEST_EFFORT_UPDATE:
            for name in BASE_FIELDS:
                if name in FIELD_LABELS and not usable_value(case, name, pin_length):
                    result[name] = None
            for prefix in ("registration_address", "living_address"):
                if not usable_value(case, prefix):
                    for name in list(result):
                        if name == prefix or name.startswith(prefix + "."):
                            result[name] = None
            result["phones"] = "; ".join(
                dict.fromkeys(
                    getattr(case.data, name)
                    for name in ("whatsapp_phone", "additional_phone")
                    if usable_value(case, name, pin_length)
                )
            )
        return result


@dataclass
class Change:
    sheet: str
    cell: str
    old: Any
    new: Any
    source: str
    warning: str = ""
    formula: bool = False
    field_name: str = ""


@dataclass
class ExportPlan:
    source: Path
    source_hash: str
    case_digest: str
    profile_digest: str
    changes: list[Change]
    new_sheets: list[tuple[str, str]] = field(default_factory=list)
    photos: PhotoExportPlan | None = None
    image_replace_confirmed: bool = False
    existing_contract: bool = False
    identity_conflict: bool = False
    missing_personal: list[str] = field(default_factory=list)
    preserved: list[Change] = field(default_factory=list)
    clear_missing_confirmed: bool = False
    preserved_confirmed: bool = False
    clear_contacts_confirmed: bool = False
    best_effort: bool = False
    warnings: RecognitionWarnings = field(default_factory=RecognitionWarnings)
    missing_value_policy: MissingValuePolicy = MissingValuePolicy.LEAVE_EMPTY
    policy_confirmed: bool = False
    formula_clear_errors: list[str] = field(default_factory=list)
    personal_formula_cells: list[tuple[str, str]] = field(default_factory=list)

    @property
    def needs_missing_policy(self) -> bool:
        return bool(
            self.missing_personal
            or self.formula_clear_errors
            or any(not (isinstance(c.old, str) and c.old.startswith("=")) for c in self.preserved)
            or self.photos
            and self.photos.preserved_ids
        )

    def select_missing_policy(self, policy: MissingValuePolicy) -> None:
        # This is one document-wide decision. The writer independently rebuilds it.
        if self.needs_missing_policy and policy == MissingValuePolicy.LEAVE_EMPTY:
            raise AppError("Выберите сохранение прежних значений для того же клиента или очистку для нового.")
        if self.identity_conflict and policy != MissingValuePolicy.CLEAR_OLD_PERSONAL:
            raise AppError("Обнаружена смена клиента. Старые персональные сведения необходимо очистить.")
        if policy == MissingValuePolicy.CLEAR_OLD_PERSONAL:
            if self.formula_clear_errors:
                raise AppError(
                    "Нельзя безопасно очистить старые данные в формуле. Скорректируйте расположение полей: "
                    + ", ".join(self.formula_clear_errors)
                )
            for item in list(self.preserved):
                if isinstance(item.old, str) and item.old.startswith("="):
                    continue
                self.changes.append(
                    Change(
                        item.sheet,
                        item.cell,
                        item.old,
                        None,
                        "Очистка прежних персональных данных",
                        field_name=item.field_name,
                    )
                )
                self.preserved.remove(item)
            if self.photos:
                self.photos.clear_missing()
        self.missing_value_policy = policy
        self.policy_confirmed = True
        self.clear_missing_confirmed = True
        self.preserved_confirmed = True

    def clear_preserved_contacts(self) -> None:
        if self.clear_contacts_confirmed:
            return
        cleared = set()
        for item in self.preserved:
            if item.field_name in CONDITIONAL_FIELDS and not (
                isinstance(item.old, str) and item.old.startswith("=")
            ):
                self.changes.append(
                    Change(
                        item.sheet,
                        item.cell,
                        item.old,
                        None,
                        "явная очистка контактов",
                        field_name=item.field_name,
                    )
                )
                cleared.add((item.sheet, item.cell))
        self.preserved = [item for item in self.preserved if (item.sheet, item.cell) not in cleared]
        self.clear_contacts_confirmed = True


PERSONAL_FIELDS = {
    "full_name",
    "surname",
    "first_name",
    "patronymic",
    "surname_cyrillic",
    "surname_latin",
    "first_name_cyrillic",
    "first_name_latin",
    "patronymic_cyrillic",
    "birth_date",
    "personal_identifier",
    "passport_type",
    "passport_number",
    "passport_issued_by",
    "passport_department_code",
    "passport_issued_date",
    "registration_address",
    "place_of_birth",
    "gender",
    "citizenship",
    "expiration_date",
} | {f"registration_address.{part}" for part in ("region", "street", "full")}
CONDITIONAL_FIELDS = {"phones", "whatsapp_phone", "additional_phone", "living_address", "object_address"} | {
    f"living_address.{part}" for part in ("region", "street", "full")
}
EXISTING_FIELDS = PERSONAL_FIELDS | CONDITIONAL_FIELDS


def workbook_structure(book, profile: TemplateProfile) -> str:
    """Conservative custom-map identity: layout AND the untouched contract.

    Personal values may change. A changed label, formula, style, price or tool
    outside the client mapping requires reviewing the saved map again.
    Dimensions/merges alone cannot distinguish unrelated contracts.
    """
    targets = {(m.sheet, m.cell) for m in profile.field_mappings if m.field in EXISTING_FIELDS}
    layout = [
        (
            sheet.title,
            sheet.max_row,
            sheet.max_column,
            sorted(str(m) for m in sheet.merged_cells.ranges),
            sorted(
                (cell.coordinate, repr(cell.value), str(cell._style), cell.data_type)
                for cell in sheet._cells.values()
                if (cell.value is not None or cell.has_style)
                and (sheet.title, cell.coordinate) not in targets
            ),
            sorted(
                (address, str(getattr(sheet._cells.get(coordinate_to_tuple(address)), "_style", None)))
                for name, address in targets
                if name == sheet.title
            ),
        )
        for sheet in book
    ]
    return hashlib.sha256(repr(layout).encode()).hexdigest()


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    try:
        with path.open("rb") as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                digest.update(chunk)
    except OSError as exc:
        raise AppError(
            "Не удалось прочитать Excel. Проверьте доступ к файлу; возможно, он занят другой программой."
        ) from exc
    return digest.hexdigest()


def profile_hash(profile: TemplateProfile) -> str:
    return hashlib.sha256(profile.model_dump_json().encode()).hexdigest()


def inspect_workbook(path: Path) -> None:
    if path.suffix.lower() != ".xlsx":
        raise AppError("Поддерживаются только XLSX. XLSM с макросами не поддерживается.")
    try:
        if path.stat().st_size > 50_000_000:
            raise AppError("Excel-файл больше 50 МБ. Используйте сокращённый шаблон.")
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > 5000 or sum(i.file_size for i in infos) > 150_000_000:
                raise AppError("Excel-файл превышает безопасные ограничения распаковки.")
            unsupported = (
                "xl/activeX/",
                "xl/embeddings/",
                "_xmlsignatures/",
            )
            if any(i.filename.startswith(unsupported) or "vbaProject" in i.filename for i in infos):
                raise AppError(
                    "Шаблон содержит макросы, подпись, ActiveX или встроенные специальные объекты. Такой XLSX не поддерживается."
                )
            for info in infos:
                if info.filename.endswith((".xml", ".rels")):
                    xml = archive.read(info).upper()
                    if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                        raise AppError("Небезопасные XML-объявления в Excel.")
                if info.file_size > 5_000_000 and info.file_size / max(1, info.compress_size) > 500:
                    raise AppError("Небезопасная структура Excel-файла.")
            if archive.testzip() is not None:
                raise AppError("Excel-файл повреждён.")
    except AppError:
        raise
    except (OSError, zipfile.BadZipFile, RuntimeError) as exc:
        raise AppError("Не удалось открыть Excel. Проверьте формат, доступ и целостность файла.") from exc


def open_workbook(path: Path):
    inspect_workbook(path)
    try:
        book = load_workbook(path, data_only=False, keep_links=True, rich_text=True)
        if any(s.max_row * s.max_column > 2_000_000 for s in book):
            book.close()
            raise AppError(
                "Область Excel слишком велика для безопасной обработки. Сократите пустое форматирование."
            )
        return book
    except AppError:
        raise
    except Exception as exc:
        # Third-party XML parsers raise several implementation-specific exceptions.
        raise AppError("Не удалось прочитать структуру Excel-шаблона.") from exc


class WorkbookService:
    def describe(
        self,
        path: Path,
        current: TemplateProfile,
        *,
        strict: bool = False,
        registry: TemplateRegistry | None = None,
    ) -> tuple[list[str], TemplateProfile]:
        book = open_workbook(path)
        try:
            recognized = (
                {"Дог.", "Пр №1", "Акт прием", "Акт расч"} <= set(book.sheetnames)
                and {"D7:I7", "A21:H21", "I21:P21"} <= {str(m) for m in book["Дог."].merged_cells.ranges}
                and {"B20:E21", "B22:E23", "J20:K20", "N28:O29"}
                <= {str(m) for m in book["Пр №1"].merged_cells.ranges}
                and book["Пр №1"]["J19"].value == "Дата и время выдачи   Имущества"
                and book["Акт прием"]["L20"].value == "Завод. или индеф. №"
            )
            previous = lermontov_profile().model_copy(
                update={
                    "version": "2",
                    "photo_sheet": None,
                    "photos_required": False,
                    "allow_photo_exception": False,
                }
            )
            presets = [legacy_profile("combined"), legacy_profile("separate"), lermontov_profile(), previous]
            if strict:
                if recognized:
                    return book.sheetnames, lermontov_profile()
                if current.structure_digest and current.structure_digest == workbook_structure(book, current):
                    return book.sheetnames, current
                saved = registry.matching(book) if registry else None
                if saved:
                    return book.sheetnames, saved
                demo = (
                    all(
                        name in book.sheetnames
                        and book[name]["A1"].value == "ДЕМОНСТРАЦИОННЫЙ ШАБЛОН — НЕ ДОГОВОР"
                        for name in ("Дог.", "Пр №1", "Акт расч", "Акт прием")
                    )
                    and "Карта полей" in book.sheetnames
                )
                if demo:
                    selected = legacy_profile(current.phone_mode)
                    for mapping in selected.field_mappings:
                        if mapping.field in {"birth_date", "passport_issued_date"}:
                            mapping.value_type = "date"
                    if "Пасп Доп Данные" in book.sheetnames:
                        selected.photo_sheet = "Пасп Доп Данные"
                        selected.photos_required = True
                        selected.allow_photo_exception = True
                    return book.sheetnames, selected
                return book.sheetnames, TemplateProfile(id="unmapped", tool_sheet=None, field_mappings=[])
            if current in presets:
                current = lermontov_profile() if recognized else legacy_profile(current.phone_mode)
            return book.sheetnames, current
        finally:
            book.close()

    def bind_profile(self, path: Path, profile: TemplateProfile) -> TemplateProfile:
        book = open_workbook(path)
        try:
            for mapping in profile.field_mappings:
                if mapping.sheet not in book.sheetnames:
                    if mapping.sheet in profile.optional_sheets:
                        continue
                    raise AppError(f"Не найден лист «{mapping.sheet}» из mapping.")
                self._change(book[mapping.sheet], mapping.cell, None, "проверка расположения", True)
            return profile.model_copy(update={"structure_digest": workbook_structure(book, profile)})
        finally:
            book.close()

    @staticmethod
    def _copy_page(book, source, name):
        if source.tables:
            raise AppError(
                "Страница инструментов содержит Excel-таблицу с уникальным именем. Используйте шаблон с обычными ячейками."
            )
        target = book.copy_worksheet(source)
        target.title = name
        # openpyxl's worksheet copier omits print areas, repeating titles and several
        # view/validation attributes; copy these explicitly for additional contract pages.
        if source.print_area:
            target.print_area = [str(r) for r in PrintArea.from_string(source.print_area).ranges]
        target.print_title_rows = source.print_title_rows
        target.print_title_cols = source.print_title_cols
        for attribute in (
            "HeaderFooter",
            "views",
            "row_breaks",
            "col_breaks",
            "auto_filter",
            "data_validations",
            "conditional_formatting",
        ):
            setattr(target, attribute, deepcopy(getattr(source, attribute)))
        target.freeze_panes = source.freeze_panes
        target.sheet_state = source.sheet_state
        return target

    def sheets(self, path: Path) -> list[str]:
        book = open_workbook(path)
        try:
            return book.sheetnames
        finally:
            book.close()

    def plan(self, source: Path, case: ClientCase, profile: TemplateProfile) -> ExportPlan:
        problems = readiness(case, profile)
        if problems:
            raise AppError("Экспорт пока невозможен:\n" + "\n".join(problems))
        source = source.resolve()
        if (
            case.workflow_mode.updates_existing
            and case.template_path
            and source != case.template_path.resolve()
        ):
            raise AppError("Выбран другой исходный договор. Повторите выбор файла и подтверждение.")
        initial_hash = file_hash(source)
        book = open_workbook(source)
        changes: list[Change] = []
        copies: list[tuple[str, str]] = []
        try:
            for sheet in profile.required_sheets:
                if sheet not in book.sheetnames:
                    raise AppError(f"В шаблоне отсутствует обязательный лист «{sheet}».")
            assembled = ExcelDataAssembler().assemble(case, profile.pin_length)
            if case.workflow_mode.updates_existing:
                return self._existing_plan(source, initial_hash, book, case, profile, assembled)
            for mapping in profile.field_mappings:
                if mapping.sheet not in book.sheetnames:
                    if mapping.sheet in profile.optional_sheets:
                        continue
                    raise AppError(f"Не найден лист «{mapping.sheet}» из mapping.")
                value = assembled.get(mapping.field)
                if value in (None, "") and mapping.skip_empty:
                    continue
                value = self._typed_value(value, mapping.value_type)
                evidence = case.evidence.get(mapping.field.split(".")[0])
                source_name = evidence.source if evidence else "ручной ввод / данные договора"
                changes.append(
                    self._change(
                        book[mapping.sheet], mapping.cell, value, source_name, mapping.allow_formula_replace
                    )
                )
            self._tools(book, case.contract, profile, changes, copies)
            for original, name in copies:
                for mapping in profile.field_mappings:
                    if mapping.sheet != original:
                        continue
                    value = assembled.get(mapping.field)
                    if value in (None, "") and mapping.skip_empty:
                        continue
                    changes.append(
                        self._change(
                            book[name],
                            mapping.cell,
                            self._typed_value(value, mapping.value_type),
                            "данные договора на дополнительной странице",
                            mapping.allow_formula_replace,
                        )
                    )
            if profile.enable_return_formula and case.contract.actual_return_date:
                if profile.id == "lermontov-combined":
                    raise AppError(
                        "В этом договоре расчёт возврата уже задан формулами. Отключите старое правило L21 в настройках профиля; дата возврата записывается в приложение."
                    )
                if "Акт расч" not in book.sheetnames:
                    raise AppError("Не найден «Акт расч» для расчёта возврата.")
                sheet = book["Акт расч"]
                # The historical subtraction is meaningful only when both operands are dates.
                if not isinstance(sheet["F21"].value, (date,)) or not sheet["F21"].is_date:
                    from datetime import datetime

                    if not isinstance(sheet["F21"].value, datetime):
                        raise AppError("Формула G21-F21 требует проверенной даты в F21.")
                for cell in ("H20", "H21"):
                    changes.append(self._change(sheet, cell, None, "правило возврата", False))
                change = self._change(sheet, "L21", "=G21-F21", "правило профиля", True)
                change.formula = True
                changes.append(change)
            if case.contract.purchase_price is not None:
                names = (
                    [profile.acceptance_sheet]
                    if profile.acceptance_sheet
                    else [n for n in profile.optional_sheets if n.startswith("Акт при")]
                )
                matches = [n for n in names if n in book.sheetnames]
                if len(matches) > 1:
                    raise AppError("Найдено несколько актов приёма. Укажите точное имя листа в профиле.")
                if matches:
                    if not any(c.sheet == matches[0] and c.cell == "N21" for c in changes):
                        changes.append(
                            self._change(
                                book[matches[0]], "N21", case.contract.purchase_price, "ручной ввод", False
                            )
                        )
            targets = [(c.sheet, c.cell) for c in changes]
            if not targets:
                raise AppError("Нет ячеек для записи. Настройте mapping и заполните данные.")
            if len(targets) != len(set(targets)):
                raise AppError("Mapping пересекается с таблицей инструментов или правилами договора.")
            if file_hash(source) != initial_hash:
                raise AppError("Шаблон изменился во время чтения. Повторите предпросмотр.")
            photos = ExcelImageEmbedder().plan(PackageEditor(source), book, case, profile)
            return ExportPlan(
                source, initial_hash, case.digest(), profile_hash(profile), changes, copies, photos
            )
        finally:
            book.close()

    def _existing_plan(self, source, initial_hash, book, case, profile, assembled) -> ExportPlan:
        _, verified = self.describe(source, profile, strict=True)
        if not verified.field_mappings:
            raise AppError(
                "Структура договора неизвестна. В разделе «Дополнительно» подтвердите расположение полей один раз."
            )
        if case.workbook_profile is not None and case.workbook_profile != profile.model_dump(mode="json"):
            raise AppError("Профиль выбранного договора изменился. Повторите выбор и подтверждение.")
        trusted = {(m.field, m.sheet, m.cell) for m in verified.field_mappings if m.field in EXISTING_FIELDS}
        if profile.photo_sheet != verified.photo_sheet:
            raise AppError("Лист фотографий не соответствует проверенному профилю договора.")
        mappings = [
            m
            for m in profile.field_mappings
            if m.field in EXISTING_FIELDS
            and (case.workflow_mode != WorkflowMode.BEST_EFFORT_UPDATE or m.field != "object_address")
        ]
        if case.workflow_mode != WorkflowMode.BEST_EFFORT_UPDATE and not {
            "full_name",
            "personal_identifier",
            "passport_number",
        } <= {m.field for m in mappings}:
            raise AppError("В профиле должны быть проверенные ячейки ФИО, ПИН и номера документа.")
        for mapping in mappings:
            if (mapping.field, mapping.sheet, mapping.cell) not in trusted:
                raise AppError("Обнаружена неразрешённая ячейка. Проверьте расположение полей договора.")

        # Name comparisons ignore case, whitespace and punctuation, not letters/digits.
        def identity(value):
            return re.sub(r"\W", "", str(value or "")).casefold()

        def differs(mapping):
            old, new = book[mapping.sheet][mapping.cell].value, assembled.get(mapping.field)
            if identity(old) == identity(new):
                return False
            if case.workflow_mode == WorkflowMode.BEST_EFFORT_UPDATE and mapping.field == "full_name":
                before = re.findall(r"[^\W\d_]+", str(old).casefold())
                after = re.findall(r"[^\W\d_]+", str(new).casefold())
                if len(before) != len(after) and (
                    before[: len(after)] == after or after[: len(before)] == before
                ):
                    return False
            return True

        conflict = any(
            m.field in {"full_name", "personal_identifier", "passport_number"}
            and book[m.sheet][m.cell].data_type != "f"
            and book[m.sheet][m.cell].value not in (None, "")
            and assembled.get(m.field) not in (None, "")
            and differs(m)
            for m in mappings
            if m.sheet in book.sheetnames
        )
        if case.workflow_mode == WorkflowMode.BEST_EFFORT_UPDATE:
            return self._best_effort_plan(
                source, initial_hash, book, case, profile, assembled, mappings, conflict
            )
        changes, preserved, missing = [], [], []
        for mapping in mappings:
            if mapping.sheet not in book.sheetnames:
                if mapping.sheet in profile.optional_sheets:
                    continue
                raise AppError(f"Не найден лист «{mapping.sheet}» из mapping.")
            cell = book[mapping.sheet][mapping.cell]
            value = assembled.get(mapping.field)
            empty = value in (None, "")
            clear = empty and conflict and mapping.field in PERSONAL_FIELDS and cell.value not in (None, "")
            label = FIELD_LABELS.get(mapping.field.split(".")[0], "Телефоны")
            if empty and not clear:
                if cell.value not in (None, ""):
                    preserved.append(
                        Change(
                            mapping.sheet,
                            mapping.cell,
                            cell.value,
                            cell.value,
                            label,
                            field_name=mapping.field,
                        )
                    )
                continue
            # Formula permissions from the template generator do not grant
            # permission in the ordinary existing-contract workflow.
            change = self._change(
                book[mapping.sheet],
                mapping.cell,
                self._typed_value(
                    value,
                    "text"
                    if mapping.field
                    in {"personal_identifier", "passport_type", "passport_number", "passport_department_code"}
                    else mapping.value_type,
                ),
                label,
                False,
            )
            change.field_name = mapping.field
            if clear:
                change.warning = (
                    "Старое персональное значение будет очищено только после явного подтверждения"
                )
                missing.append(label)
            if change.old != change.new:
                changes.append(change)
        if not changes:
            # A repeat save can still replace the client's photographs.
            if not profile.photo_sheet:
                raise AppError("Нет изменений данных клиента для сохранения.")
        photos = ExcelImageEmbedder().plan(PackageEditor(source), book, case, profile)
        if file_hash(source) != initial_hash:
            raise AppError("Исходный Excel изменился во время чтения. Повторите предпросмотр.")
        return ExportPlan(
            source,
            initial_hash,
            case.digest(),
            profile_hash(profile),
            changes,
            photos=photos,
            existing_contract=True,
            identity_conflict=conflict,
            missing_personal=missing,
            preserved=preserved,
        )

    def _best_effort_plan(self, source, initial_hash, book, case, profile, assembled, mappings, conflict):
        changes, preserved, missing, formula_errors = [], [], [], []
        targets = {(m.sheet, m.cell) for m in mappings}
        formula_cells = []

        def safe_reference(sheet, address, seen=None):
            seen = set() if seen is None else seen
            key = (sheet, address)
            if key in seen or key not in targets:
                return False
            seen.add(key)
            cell = book[sheet][address]
            if cell.data_type != "f":
                return True
            ref = re.fullmatch(r"=\$?([A-Z]+)\$?([0-9]+)", str(cell.value))
            return bool(ref and safe_reference(sheet, ref[1] + ref[2], seen))

        for mapping in mappings:
            if mapping.sheet not in book.sheetnames:
                if mapping.sheet in profile.optional_sheets:
                    continue
                raise AppError(f"Не найден лист «{mapping.sheet}» из карты полей.")
            sheet = book[mapping.sheet]
            cell = sheet[mapping.cell]
            value = assembled.get(mapping.field)
            base = mapping.field.split(".")[0]
            if base in FIELD_LABELS and not usable_value(case, base, profile.pin_length):
                value = None
            label = FIELD_LABELS.get(base, "Телефоны")
            if value in (None, ""):
                if cell.value not in (None, ""):
                    item = Change(
                        mapping.sheet, mapping.cell, cell.value, cell.value, label, field_name=mapping.field
                    )
                    if cell.data_type == "f":
                        # A missing value never grants permission to erase a formula.
                        # Only a direct reference into the mapped client area can follow a client change safely.
                        formula_cells.append((mapping.sheet, mapping.cell))
                        if not safe_reference(mapping.sheet, mapping.cell):
                            formula_errors.append(label)
                        preserved.append(item)
                    elif conflict and mapping.field in PERSONAL_FIELDS:
                        change = self._change(
                            sheet, mapping.cell, None, "Очистка данных прежнего клиента", False
                        )
                        change.field_name = mapping.field
                        changes.append(change)
                        missing.append(label)
                    else:
                        preserved.append(item)
                continue
            change = self._change(
                sheet,
                mapping.cell,
                self._typed_value(
                    value,
                    "text"
                    if mapping.field
                    in {"personal_identifier", "passport_type", "passport_number", "passport_department_code"}
                    else mapping.value_type,
                ),
                label,
                False,
            )
            change.field_name = mapping.field
            if change.old != change.new:
                changes.append(change)
        photos = ExcelImageEmbedder().plan(PackageEditor(source), book, case, profile)
        if file_hash(source) != initial_hash:
            raise AppError("Исходный Excel изменился во время чтения. Повторите сохранение.")
        plan = ExportPlan(
            source,
            initial_hash,
            case.digest(),
            profile_hash(profile),
            changes,
            photos=photos,
            existing_contract=True,
            identity_conflict=conflict,
            missing_personal=missing,
            preserved=preserved,
            best_effort=True,
            warnings=recognition_warnings(case, profile),
            formula_clear_errors=formula_errors,
            personal_formula_cells=formula_cells,
        )
        if conflict:
            # Missing fields of a known different client cannot be retained.
            plan.select_missing_policy(MissingValuePolicy.CLEAR_OLD_PERSONAL)
            plan.policy_confirmed = False  # The single final summary still needs acceptance.
        return plan

    @staticmethod
    def _typed_value(value: Any, kind: str) -> Any:
        if value in (None, ""):
            return None
        if kind == "text":
            return value.strftime("%d.%m.%Y") if isinstance(value, date) else str(value)
        if kind == "date":
            from document_excel.domain.models import parse_date

            try:
                return value if isinstance(value, date) else parse_date(str(value))
            except ValueError as exc:
                raise AppError("В mapping указана дата для недатового значения.") from exc
        try:
            number = float(value)
            if not __import__("math").isfinite(number):
                raise ValueError
            return number
        except (ValueError, TypeError) as exc:
            raise AppError("В mapping указано число для нечислового значения.") from exc

    @staticmethod
    def _change(sheet, address: str, value: Any, source: str, allow_formula: bool) -> Change:
        row, column = coordinate_to_tuple(address)
        if row > sheet.max_row or column > sheet.max_column:
            raise AppError(f"Ячейка {sheet.title}!{address} вне существующей области шаблона.")
        if sheet.protection.sheet:
            raise AppError(f"Лист «{sheet.title}» защищён. Используйте разрешённую копию шаблона.")
        cell = sheet[address]
        if isinstance(cell, MergedCell):
            raise AppError(f"{sheet.title}!{address} — неведущая ячейка объединения. Исправьте mapping.")
        if cell.data_type == "f" and not allow_formula:
            raise AppError(f"{sheet.title}!{address} содержит формулу. Замена запрещена профилем.")
        warning = (
            "Будет заменено существующее значение" if cell.value is not None and cell.value != value else ""
        )
        if cell.data_type == "f":
            warning = "Замена формулы явно разрешена профилем"
        return Change(sheet.title, address, cell.value, value, source, warning)

    def _tools(
        self,
        book,
        contract: ContractData,
        profile: TemplateProfile,
        changes: list[Change],
        copies: list[tuple[str, str]],
    ) -> None:
        if not contract.tools and contract.deposit_amount is None:
            return
        if not profile.tool_sheet or profile.tool_sheet not in book.sheetnames:
            raise AppError("Для инструментов / залога требуется настроенный лист приложения.")
        base = book[profile.tool_sheet]
        page_count = max(1, (len(contract.tools) + profile.tools_per_page - 1) // profile.tools_per_page)
        # Existing historical pages would retain stale clients; require explicit cleanup in the source copy.
        stale = [name for name in book.sheetnames if re.fullmatch(re.escape(base.title) + r" \(\d+\)", name)]
        if stale:
            raise AppError(
                "В исходнике уже есть дополнительные страницы инструментов. Используйте чистый шаблон без них."
            )
        for page in range(page_count):
            sheet = base
            linked = {name: book[name] for name in profile.tool_linked_sheets if name in book.sheetnames}
            if page:
                names = {name: f"{name} ({page + 1})" for name in [base.title, *linked]}
                for original, name in names.items():
                    if len(name) > 31 or name in book.sheetnames:
                        raise AppError("Имя дополнительного листа занято или превышает 31 символ.")
                    target = self._copy_page(book, book[original], name)
                    copies.append((original, name))
                    for row in target:
                        for cell in row:
                            if cell.data_type == "f" and isinstance(cell.value, str):
                                cell.value = remap_formula(cell.value, names)
                    if original == base.title:
                        sheet = target
                    else:
                        linked[original] = target
            for offset in range(profile.tools_per_page):
                position = page * profile.tools_per_page + offset
                tool = contract.tools[position] if position < len(contract.tools) else None
                values = {
                    "A": position + 1 if tool else None,
                    "B": tool.name if tool else None,
                    "I": tool.quantity if tool else None,
                    "N": tool.price_per_day if tool else None,
                }
                if profile.tool_inventory_column:
                    values[profile.tool_inventory_column] = tool.inventory_number if tool else None
                for column, value in values.items():
                    changes.append(
                        self._change(
                            sheet,
                            f"{column}{profile.tool_first_row + offset * profile.tool_row_step}",
                            value,
                            "ручная таблица инструментов",
                            column == "B" and profile.tool_name_formula_replace,
                        )
                    )
                if profile.tool_inventory_sheet:
                    inventory_sheet = linked.get(profile.tool_inventory_sheet)
                    if inventory_sheet is None:
                        raise AppError("Не найден лист заводских номеров инструментов.")
                    changes.append(
                        self._change(
                            inventory_sheet,
                            f"{profile.tool_inventory_target_column}{profile.tool_inventory_first_row + offset}",
                            tool.inventory_number if tool else None,
                            "заводской номер инструмента",
                            False,
                        )
                    )
            changes.append(
                self._change(
                    sheet,
                    "P20",
                    contract.deposit_amount if page == 0 else None,
                    "залог первой страницы",
                    False,
                )
            )

    def update_existing(
        self, plan: ExportPlan, case: ClientCase, profile: TemplateProfile, *, new_name: str | None = None
    ) -> Path:
        """Explicitly replace the selected workbook after validating a temporary result."""
        if (
            not case.workflow_mode.updates_existing
            or not case.template_path
            or case.template_path.resolve() != plan.source
        ):
            raise AppError("Для обновления выберите тот же исходный Excel-договор.")
        return self.write(plan, plan.source, case, profile, replace_source=True, rename_to=new_name)

    def write(
        self,
        plan: ExportPlan,
        destination: Path,
        case: ClientCase,
        profile: TemplateProfile,
        *,
        replace_source: bool = False,
        rename_to: str | None = None,
    ) -> Path:
        if (
            readiness(case, profile)
            or case.digest() != plan.case_digest
            or profile_hash(profile) != plan.profile_digest
        ):
            raise AppError("Данные или профиль изменились. Повторите подтверждение и предпросмотр.")
        if rename_to is not None:
            if (
                not replace_source
                or not rename_to
                or Path(rename_to).name != rename_to
                or re.search(r'[<>:"/\\|?*\x00-\x1f]', rename_to)
                or rename_to.rstrip(" .") != rename_to
                or Path(rename_to).suffix.lower() != ".xlsx"
                or re.fullmatch(r"(?i:CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\..*)?", rename_to)
            ):
                raise AppError("Укажите безопасное имя XLSX в той же папке выбранного договора.")
        destination = destination.resolve()
        if destination.suffix.lower() != ".xlsx":
            raise AppError("Сохраните файл с расширением .xlsx.")
        if replace_source:
            if destination != plan.source or not case.workflow_mode.updates_existing:
                raise AppError("Разрешено обновлять только выбранный Excel-договор.")
        elif destination == plan.source:
            raise AppError("Для изменения выбранного договора используйте явное обновление документа.")
        if replace_source:
            try:
                read_only = not (plan.source.stat().st_mode & stat.S_IWUSR)
            except OSError as exc:
                raise AppError("Выбранный Excel недоступен. Проверьте файл и повторите сохранение.") from exc
            if read_only:
                raise AppError("Выбранный Excel доступен только для чтения. Разрешите запись в этот файл.")
        if destination.exists() and not replace_source:
            raise OutputExistsError("Файл уже существует. Выберите другое имя; перезапись отключена.")
        if file_hash(plan.source) != plan.source_hash:
            raise AppError("Исходный Excel изменился. Повторите предпросмотр.")
        if case.workflow_mode.updates_existing:
            expected = self.plan(plan.source, case, profile)
            if plan.best_effort and plan.policy_confirmed:
                expected.select_missing_policy(plan.missing_value_policy)
            if plan.clear_contacts_confirmed:
                expected.clear_preserved_contacts()
            if (
                not plan.existing_contract
                or plan.best_effort != expected.best_effort
                or plan.changes != expected.changes
                or plan.new_sheets
                or plan.photos != expected.photos
                or plan.missing_personal != expected.missing_personal
                or plan.preserved != expected.preserved
                or plan.identity_conflict != expected.identity_conflict
                or plan.personal_formula_cells != expected.personal_formula_cells
            ):
                raise AppError(
                    "Обнаружено неожиданное изменение: разрешены только проверенные поля клиента и фотографии."
                )
            if (
                plan.best_effort
                and (plan.needs_missing_policy or plan.identity_conflict)
                and not plan.policy_confirmed
            ):
                raise AppError("Выберите общее действие со старыми персональными данными в итоговой сводке.")
            if plan.missing_personal and not plan.clear_missing_confirmed:
                raise AppError(
                    "Договор содержит данные другого клиента. Заполните отсутствующие поля или явно подтвердите их очистку."
                )
            if plan.preserved and not plan.preserved_confirmed:
                raise AppError(
                    "Проверьте сохраняемые значения исходного договора и подтвердите их принадлежность текущему клиенту."
                )
        book = open_workbook(plan.source)
        temporary: Path | None = None
        try:
            editor = PackageEditor(plan.source)
            for original, name in plan.new_sheets:
                # Related pages with the same suffix must link to one another.
                suffix = name[len(original) :]
                names = {o: n for o, n in plan.new_sheets if n[len(o) :] == suffix}
                target = self._copy_page(book, book[original], name)
                for row in target:
                    for cell in row:
                        if cell.data_type == "f" and isinstance(cell.value, str):
                            cell.value = remap_formula(cell.value, names)
                editor.clone_sheet(original, name, names)
            allowed = {(change.sheet, change.cell) for change in plan.changes}
            untouched = self._snapshot(book, allowed)
            for sheet in dict.fromkeys(c.sheet for c in plan.changes):
                editor.cells(sheet, [c for c in plan.changes if c.sheet == sheet], book)
            if plan.existing_contract and plan.changes:
                editor.scrub_personal_shared_strings()
                editor.clear_formula_caches(plan.personal_formula_cells)
            if plan.photos is not None:
                ExcelImageEmbedder().apply(editor, plan.photos, plan.image_replace_confirmed)
            destination.parent.mkdir(parents=True, exist_ok=True)
            fd, name = tempfile.mkstemp(prefix=".document-excel-", suffix=".xlsx", dir=destination.parent)
            os.close(fd)
            temporary = Path(name)
            if (
                not plan.changes
                and not plan.new_sheets
                and not (plan.photos and (plan.photos.insertions or plan.photos.remove_ids))
            ):
                shutil.copyfile(plan.source, temporary)
            else:
                editor.save(temporary)
            if plan.photos is not None:
                ExcelImageEmbedder().verify(PackageEditor(temporary), plan.photos)
            check = open_workbook(temporary)
            try:
                if self._snapshot(check, allowed) != untouched:
                    raise AppError("Проверка сохранности Excel не пройдена. Результат не сохранён.")
                for change in plan.changes:
                    cell = check[change.sheet][change.cell]
                    expected = change.new
                    actual = cell.value
                    if isinstance(expected, date) and hasattr(actual, "date"):
                        actual = actual.date()
                    if expected == "" and actual is None:
                        continue
                    if actual != expected or (
                        isinstance(expected, str) and not change.formula and cell.data_type == "f"
                    ):
                        raise AppError("Проверка записанных значений Excel не пройдена.")
            finally:
                check.close()
            if file_hash(plan.source) != plan.source_hash:
                raise AppError("Исходный файл изменился во время сохранения. Повторите экспорт.")
            with temporary.open("r+b") as stream:
                os.fsync(stream.fileno())
            if replace_source:
                # All readers are closed before replacing the selected file on Windows.
                book.close()
                shutil.copymode(plan.source, temporary)
                if file_hash(plan.source) != plan.source_hash:
                    raise AppError("Выбранный Excel изменился перед записью. Повторите сохранение.")
                moved = False
                if rename_to:
                    base = plan.source.with_name(rename_to)
                    for _ in range(64):
                        destination = ContractFilenameGenerator.available(base, current=plan.source)
                        if destination == plan.source:
                            break
                        try:
                            rename_without_replace(plan.source, destination)
                            moved = True
                            break
                        except FileExistsError:
                            continue
                    else:
                        raise AppError("Не удалось подобрать свободное имя договора. Повторите сохранение.")
                try:
                    os.replace(temporary, destination)
                except OSError as exc:
                    if moved:
                        try:
                            if file_hash(destination) != plan.source_hash:
                                raise AppError("Файл изменился после переименования.")
                            rename_without_replace(destination, plan.source)
                        except (OSError, AppError) as rollback_error:
                            raise AppError(
                                "Сохранение прервано. Файл доступен по адресу: " + str(destination)
                            ) from rollback_error
                    raise exc
                temporary = None
            else:
                publish_copy(temporary, destination)
                temporary.unlink()
                temporary = None
            return destination
        except AppError:
            raise
        except FileExistsError as exc:
            raise OutputExistsError(
                "Файл уже существует. Выберите другое имя или сохраните с дополнительным номером."
            ) from exc
        except OSError as exc:
            raise AppError(
                "Не удалось обновить выбранный Excel. Закройте его в Excel, проверьте права и свободное место и повторите сохранение."
                if replace_source
                else "Не удалось сохранить Excel. Проверьте свободное место, права и имя файла. Выберите другую папку или закройте файл, блокирующий запись."
            ) from exc
        finally:
            book.close()
            if temporary is not None:
                temporary.unlink(missing_ok=True)

    @staticmethod
    def _snapshot(book, allowed: set[tuple[str, str]]) -> str:
        payload = []
        for sheet in book:
            cells = [
                (c.coordinate, c.value, c.data_type, str(c._style), c.number_format)
                for row in sheet
                for c in row
                if (sheet.title, c.coordinate) not in allowed and (c.value is not None or c.has_style)
            ]
            payload.append(
                (
                    sheet.title,
                    cells,
                    [str(m) for m in sheet.merged_cells.ranges],
                    str(sheet.page_setup),
                    str(sheet.print_options),
                    str(sheet.page_margins),
                    str(sheet.print_area),
                    list(sheet.print_title_rows or ""),
                )
            )
        return hashlib.sha256(repr(payload).encode()).hexdigest()


def rename_without_replace(source: Path, destination: Path) -> None:
    """Move a file within its folder without replacing an occupied name or copying bytes."""
    if os.name == "nt":
        os.rename(source, destination)  # Windows refuses an existing destination.
        return
    # POSIX rename would overwrite an existing file. Exclusive hard-link creation
    # followed by unlink moves the same inode and never copies document bytes.
    os.link(source, destination)
    try:
        source.unlink()
    except OSError:
        destination.unlink()
        raise


def publish_copy(temporary: Path, destination: Path) -> None:
    """Publish without replacing an existing name, including on exFAT/FAT.

    Hard-link publication is atomic. The exclusive-copy fallback is not atomic:
    readers may see an incomplete file during copying; failure removes our file.
    """
    try:
        os.link(temporary, destination)
        return
    except FileExistsError:
        raise
    except OSError as exc:
        if exc.errno not in {
            errno.EXDEV,
            errno.EPERM,
            errno.EOPNOTSUPP,
            errno.ENOSYS,
            errno.EINVAL,
        } and getattr(exc, "winerror", None) not in {1, 50}:
            raise
    created = False
    try:
        with destination.open("xb") as target:
            created = True
            with temporary.open("rb") as source:
                shutil.copyfileobj(source, target)
            target.flush()
            os.fsync(target.fileno())
        if file_hash(temporary) != file_hash(destination):
            raise AppError("Проверка сохранённой копии не пройдена.")
    except BaseException:
        if created:
            destination.unlink(missing_ok=True)
        raise


def create_demo_template(path: Path) -> None:
    """A visibly marked synthetic workbook; never an original contract."""
    book = Workbook()
    book.remove(book.worksheets[0])
    for name in ("Дог.", "Пр №1", "Акт расч", "Акт прием"):
        sheet = book.create_sheet(name)
        for row in sheet.iter_rows(min_row=1, max_row=30, min_col=1, max_col=16):
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet["A1"] = "ДЕМОНСТРАЦИОННЫЙ ШАБЛОН — НЕ ДОГОВОР"
        sheet["A1"].font = Font(bold=True, size=15, color="B42318")
        sheet.merge_cells("A1:P2")
        for letter in "ABCDEFGHIJKLMNOP":
            sheet.column_dimensions[letter].width = 12
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.print_area = "A1:P30"
        sheet["P30"] = "Синтетический пример"
    for mapping in legacy_profile().field_mappings:
        sheet = book[mapping.sheet]
        cell = sheet[mapping.cell]
        cell.fill = PatternFill("solid", fgColor="E4EDFA")
        cell.font = Font(size=11, color="182B49")
    guide = book.create_sheet("Карта полей")
    guide.append(["ДЕМОНСТРАЦИЯ — исходный договор не предоставлен", "Лист", "Ячейка"])
    for mapping in legacy_profile().field_mappings:
        guide.append([mapping.field, mapping.sheet, mapping.cell])
    guide.column_dimensions["A"].width = 65
    guide.column_dimensions["B"].width = 20
    guide.column_dimensions["C"].width = 15
    book["Акт расч"]["F21"] = date(2026, 1, 1)
    book["Акт расч"]["F21"].number_format = "dd.mm.yyyy"
    book["Акт расч"]["L21"] = "=G21-F21"
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        book.close()
        raise AppError("Демонстрационный файл уже существует. Выберите новое имя.")
    book.save(path)
    book.close()
