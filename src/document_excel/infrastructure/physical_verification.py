"""Synthetic plastic TD1 fixtures and an end-to-end installed-package check."""

import threading
from pathlib import Path
from typing import Any, cast

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

from document_excel.application.workflows import RecognitionService
from document_excel.domain.errors import AppError
from document_excel.domain.models import (
    Address,
    ClientCase,
    FieldEvidence,
    FieldStatus,
    RegistrationStatus,
    SourceImage,
    new_id,
)
from document_excel.infrastructure.excel import WorkbookService, create_demo_template, legacy_profile
from document_excel.infrastructure.extraction import apply_candidates
from document_excel.infrastructure.mrz import check_digit
from document_excel.infrastructure.storage import Settings


def synthetic_plastic_images(output: Path, controls: dict | None = None) -> list[Path]:
    from datetime import date

    values = {
        "surname_cyrillic": "СИНТЕТИКОВ",
        "surname_latin": "SINTETIKOV",
        "first_name_cyrillic": "ТЕСТ",
        "first_name_latin": "TEST",
        "patronymic_cyrillic": "ПРИМЕРОВИЧ",
        "passport_number": "0000001",
        "personal_identifier": "00123456789012",
        "birth_date": "1990-11-07",
        "passport_issued_date": "2024-11-22",
        "expiration_date": "2034-11-22",
        "passport_issued_by": "ТЕСТ ОРГАН",
    } | (controls or {})

    def printed_date(name):
        return date.fromisoformat(values[name]).strftime("%d %m %Y")

    output.mkdir(parents=True, exist_ok=True)
    regular_path = next(
        (
            p
            for p in [
                Path("C:/Windows/Fonts/arial.ttf"),
                Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            ]
            if p.exists()
        ),
        None,
    )
    mono_path = next(
        (
            p
            for p in [
                Path("C:/Windows/Fonts/cour.ttf"),
                Path("/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"),
            ]
            if p.exists()
        ),
        None,
    )
    if regular_path is None or mono_path is None:
        raise AppError("Для синтетической проверки нужны обычный и моноширинный шрифты системы.")
    w, h = 1900, 1200

    def base():
        image = Image.new("RGB", (w, h), (228, 240, 245))
        draw = ImageDraw.Draw(image)
        draw.rounded_rectangle((15, 15, w - 15, h - 15), radius=35, outline=(60, 90, 110), width=4)
        return image, draw

    def text(draw, x, y, value, size=34, mono=False):
        font = ImageFont.truetype(str(mono_path if mono else regular_path), size)
        draw.text((round(x * w), round(y * h)), value, font=font, fill=(12, 25, 35))

    front, draw = base()
    text(draw, 0.15, 0.06, "КЫРГЫЗ РЕСПУБЛИКАСЫ / KYRGYZ REPUBLIC", 38)
    text(draw, 0.35, 0.12, "IDENTITY CARD", 38)
    draw.rectangle((100, 250, 600, 1050), fill=(165, 180, 190))
    for name, y, cyr, latin in [
        ("Фамилия / Surname", 0.225, values["surname_cyrillic"], values["surname_latin"]),
        ("Имя / Name", 0.345, values["first_name_cyrillic"], values["first_name_latin"]),
    ]:
        text(draw, 0.40, y, name, 24)
        text(draw, 0.40, y + 0.045, cyr, 38)
        text(draw, 0.40, y + 0.085, latin, 38)
    for label, value, x, y in [
        ("Отчество / Patronymic", values["patronymic_cyrillic"], 0.40, 0.47),
        ("Пол / Sex", "Э / М", 0.40, 0.55),
        ("Гражданство / Nationality", "КЫРГЫЗ РЕСПУБЛИКАСЫ", 0.40, 0.63),
        ("Дата рождения / Date of birth", printed_date("birth_date"), 0.40, 0.71),
        ("№ документа / Document No", "ID" + values["passport_number"], 0.68, 0.71),
        ("Дата окончания / Date of expiry", printed_date("expiration_date"), 0.68, 0.83),
    ]:
        text(draw, x, y, label, 21)
        if value:
            text(draw, x, y + 0.045, value, 34)
    back, draw = base()
    for label, value, x, y in [
        ("Место рождения / Place of birth", "ТЕСТОВАЯ СТРАНА", 0.27, 0.07),
        ("Орган выдачи / Authority", values["passport_issued_by"], 0.27, 0.275),
        ("Дата выдачи / Date of issue", printed_date("passport_issued_date"), 0.27, 0.35),
        ("Персональный номер / Personal number", values["personal_identifier"], 0.65, 0.20),
        ("Document No", "ID" + values["passport_number"], 0.70, 0.37),
    ]:
        text(draw, x, y, label, 20)
        text(draw, x, y + 0.045, value, 36)
    number = "ID" + values["passport_number"]
    birth = date.fromisoformat(values["birth_date"]).strftime("%y%m%d")
    expiry = date.fromisoformat(values["expiration_date"]).strftime("%y%m%d")
    a = "I<KGZ" + number + check_digit(number) + values["personal_identifier"] + "<"
    b = birth + check_digit(birth) + "M" + expiry + check_digit(expiry) + "KGZ" + "<" * 11
    b += check_digit(a[5:] + b[:7] + b[8:15] + b[18:29])
    names = (values["surname_latin"] + "<<" + values["first_name_latin"]).replace("-", "<").replace(" ", "<")
    for index, line in enumerate((a, b, names[:30].ljust(30, "<"))):
        text(draw, 0.04, 0.70 + index * 0.09, line, 94, True)
    paths = [output / "synthetic-plastic-front.png", output / "synthetic-plastic-back.png"]
    for image, path in zip((front, back), paths, strict=True):
        image.save(path)
    return paths


def run_physical_self_test(output: Path, settings: Settings | None = None) -> dict:
    images = [SourceImage(path=p) for p in synthetic_plastic_images(output)]
    result = RecognitionService(settings or Settings()).run(images, threading.Event(), lambda *_: None)
    case = ClientCase(images=images, mrz_reports=result.mrz_reports)
    for source in images:
        source.role = result.sides.get(source.id, source.role)
    apply_candidates(case, result.candidates, result.document_type, result.label_seen)
    checked = {
        "surname_cyrillic": case.data.surname_cyrillic == "СИНТЕТИКОВ",
        "surname_latin": case.data.surname_latin == "SINTETIKOV",
        "first_name_cyrillic": case.data.first_name_cyrillic == "ТЕСТ",
        "first_name_latin": case.data.first_name_latin == "TEST",
        "number": case.data.passport_number == "0000001",
        "pin": case.data.personal_identifier == "00123456789012",
        "birth_date": str(case.data.birth_date) == "1990-11-07",
        "mrz": any(r.valid for r in result.mrz_reports.values()),
        "registration_not_present": case.registration_status == RegistrationStatus.NOT_PRESENT,
    }
    if not all(checked.values()):
        return {
            "ok": False,
            "synthetic_only": True,
            "checks": checked,
            "fields_found": sorted(case.evidence),
            "conflicts": [name for name, e in case.evidence.items() if e.status == FieldStatus.CONFLICT],
        }
    case.data.registration_address = Address(region="Тестовая область", street="ул. Пример, 1")
    case.evidence["registration_address"] = FieldEvidence(
        source="ручной ввод", status=FieldStatus.MANUALLY_CORRECTED
    )
    case.registration_status = RegistrationStatus.MANUALLY_ENTERED
    template = output / "synthetic-plastic-template.xlsx"
    if not template.exists():
        create_demo_template(template)
        book = load_workbook(template)
        book.create_sheet("Пасп Доп Данные")
        book.save(template)
        book.close()
    profile = legacy_profile()
    profile.photo_sheet = "Пасп Доп Данные"
    profile.photos_required = True
    case.confirmed_digest = case.digest()
    target = output / f"synthetic-plastic-result-{new_id()[:8]}.xlsx"
    service = WorkbookService()
    plan = service.plan(template, case, profile)
    service.write(plan, target, case, profile)
    book = load_workbook(target)
    embedded = len(cast(Any, book["Пасп Доп Данные"])._images)
    book.close()
    return {
        "ok": all(checked.values()) and embedded == 2,
        "synthetic_only": True,
        "checks": checked,
        "embedded_document_photos": embedded,
        "result": target.name,
        "ocr_seconds": round(result.elapsed, 3),
    }
