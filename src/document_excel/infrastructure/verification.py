"""Packaged integration check using exclusively generated synthetic data."""

from __future__ import annotations

import threading
from importlib.resources import files
from pathlib import Path
from typing import Any, cast
from uuid import uuid4

import numpy as np
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

from document_excel.application.workflows import RecognitionService
from document_excel.domain.errors import AppError
from document_excel.domain.models import (
    ClientCase,
    DocumentSide,
    RegistrationStatus,
    SourceImage,
    WorkflowMode,
    split_document_identifier,
)
from document_excel.infrastructure.excel import WorkbookService, legacy_profile
from document_excel.infrastructure.extraction import apply_candidates
from document_excel.infrastructure.ocr_factory import create_engine
from document_excel.infrastructure.storage import Settings


def run_self_test(directory: Path) -> dict[str, object]:
    directory.mkdir(parents=True, exist_ok=True)
    resources = files("document_excel").joinpath("resources")
    image = Image.new("RGB", (1600, 1100), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(resources.joinpath("fonts/NotoSans-Regular.ttf")), 36)
    for index, line in enumerate(
        [
            "KYRGYZ REPUBLIC PASSPORT",
            "FULL NAME: SYNTHETIC TEST",
            "SURNAME: TEST",
            "GIVEN NAME: SYNTHETIC",
            "DATE OF BIRTH: 01.02.1990",
            "PERSONAL NUMBER: 00123456789012",
            "PASSPORT NO: AB1234567",
            "AUTHORITY: MKK000123",
            "REGISTRATION ADDRESS: TEST REGION, TEST STREET 1",
        ]
    ):
        draw.text((65, 65 + index * 100), line, font=font, fill="black")
    path = directory / "synthetic.png"
    image.save(path)
    source = SourceImage(path=path)
    output = RecognitionService(Settings(languages=["eng"], preprocessing="original")).run(
        [source], threading.Event(), lambda *_args: None
    )
    case = ClientCase(workflow_mode=WorkflowMode.TEMPLATE_GENERATION, images=[source])
    apply_candidates(case, output.candidates, output.document_type, output.label_seen)
    expected = {
        "full_name": "SYNTHETIC TEST",
        "personal_identifier": "00123456789012",
        "passport_type": "AB",
        "passport_number": "1234567",
        "passport_issued_by": "MKK",
        "passport_department_code": "000123",
    }
    if any(getattr(case.data, key) != value for key, value in expected.items()):
        raise AppError("Синтетическая OCR-проверка не прошла: значения отличаются от фикстуры.")
    identifier_image = Image.new("RGB", (300, 65), "white")
    ImageDraw.Draw(identifier_image).text((8, 8), "ID1234567", font=font, fill="black")
    identifier_engine = create_engine(Settings(languages=["eng"]))
    identifier_reader = identifier_engine.recognize_identifier_line
    identifier_result = identifier_reader(
        np.asarray(identifier_image), "synthetic-identifier", "original", threading.Event()
    )
    if split_document_identifier("".join(identifier_result.text.split())) != ("ID", "1234567"):
        raise AppError("Синтетическая проверка отдельной строки ID не прошла.")
    # Review is simulated exclusively for a known generated fixture, never for user documents.
    case.registration_status = RegistrationStatus.USER_CONFIRMED
    service = WorkbookService()
    template = Path(str(resources.joinpath("templates/DEMO-not-a-contract.xlsx")))
    book = load_workbook(template)
    book.create_sheet("Пасп Доп Данные")
    template = directory / f"synthetic-template-{uuid4().hex[:8]}.xlsx"
    book.save(template)
    book.close()
    source.role = DocumentSide.BACK
    front = Image.new("RGB", (1200, 750), "#dbeafe")
    ImageDraw.Draw(front).text((50, 70), "SYNTHETIC FRONT / NOT A DOCUMENT", font=font, fill="black")
    front_path = directory / "synthetic-front.png"
    front.save(front_path)
    case.images.append(SourceImage(path=front_path, role=DocumentSide.FRONT))
    case.confirmed_digest = case.digest()
    profile = legacy_profile()
    profile.photo_sheet = "Пасп Доп Данные"
    profile.photos_required = True
    plan = service.plan(template, case, profile)
    target = service.write(plan, directory / f"synthetic-result-{uuid4().hex[:8]}.xlsx", case, profile)
    check = load_workbook(target)
    try:
        if len(cast(Any, check["Пасп Доп Данные"])._images) != 2:
            raise AppError("Синтетическая проверка встроенных фотографий не прошла.")
    finally:
        check.close()
    return {
        "ok": True,
        "synthetic_only": True,
        "checked_fields": list(expected),
        "excel_reopened": True,
        "identifier_line_read": True,
        "embedded_document_photos": 2,
        "hardware_camera_tested": False,
        "result": target.name,
        "ocr_seconds": round(output.elapsed, 3),
    }
