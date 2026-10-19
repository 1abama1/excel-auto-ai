from datetime import date
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image
from PIL import Image as PillowImage

from document_excel.domain.errors import AppError
from document_excel.domain.models import Address, ContractData, ToolItem
from document_excel.infrastructure.excel import WorkbookService, file_hash, legacy_profile, lermontov_profile
from document_excel.infrastructure.xlsx_package import MAIN, document


def parts(path):
    with ZipFile(path) as archive:
        return {n: archive.read(n) for n in archive.namelist()}


def test_images_charts_external_metadata_are_preserved(template, case, profile, tmp_path):
    image = tmp_path / "logo.png"
    PillowImage.new("RGB", (20, 20), "green").save(image)
    book = load_workbook(template)
    book["Пр №1"].add_image(Image(image), "M1")
    chart = BarChart()
    book["Дог."]["N28"] = 1
    book["Дог."]["N29"] = 2
    chart.add_data(Reference(book["Дог."], min_col=14, min_row=28, max_row=29))
    book["Дог."].add_chart(chart, "Q1")
    book.save(template)
    book.close()
    # An unrelated custom part must survive even if openpyxl does not model it.
    with ZipFile(template, "a", ZIP_DEFLATED) as archive:
        archive.writestr("customXml/item1.xml", '<root xmlns="urn:test"><value>keep</value></root>')
    source = parts(template)
    case.contract = ContractData(tools=[ToolItem(name=f"Tool {i}") for i in range(7)])
    case.confirmed_digest = case.digest()
    service = WorkbookService()
    result = service.write(service.plan(template, case, profile), tmp_path / "result.xlsx", case, profile)
    output = parts(result)
    for name in source:
        if name.startswith(("xl/drawings/", "xl/charts/", "xl/media/", "customXml/")):
            assert output[name] == source[name]
    book = load_workbook(result)
    assert len(book["Дог."]._charts) == 1
    assert len(book["Пр №1"]._images) == len(book["Пр №1 (2)"]._images) == 1
    book.close()


@pytest.mark.parametrize("value", ["=1+1", "+123", "-001", "@SUM(A1)", "  текст & <пример>  "])
def test_inline_strings_are_literal(template, case, profile, tmp_path, value):
    case.data.full_name = value
    case.confirmed_digest = case.digest()
    service = WorkbookService()
    result = service.write(service.plan(template, case, profile), tmp_path / "text.xlsx", case, profile)
    book = load_workbook(result)
    assert book["Дог."]["A18"].value == value
    assert book["Дог."]["A18"].data_type == "s"
    book.close()


@pytest.mark.parametrize("bad", ["bad\x01text", "a" * 32768], ids=["control", "length"])
def test_invalid_cell_text_is_rejected(template, case, profile, tmp_path, bad):
    case.data.full_name = bad
    case.confirmed_digest = case.digest()
    target = tmp_path / "bad.xlsx"
    with pytest.raises(AppError):
        service = WorkbookService()
        service.write(service.plan(template, case, profile), target, case, profile)
    assert not target.exists()


REAL_TEMPLATE = Path(__file__).resolve().parents[1] / "assets/lermontov.xlsx"


@pytest.mark.skipif(not REAL_TEMPLATE.exists(), reason="Local user template is not distributed")
@pytest.mark.parametrize("tool_count", [0, 1, 7, 12])
def test_supplied_workbook_lossless_export(case, tmp_path, tool_count):
    service = WorkbookService()
    before = file_hash(REAL_TEMPLATE)
    # This regression covers cells/copies without photos; exercise the explicit
    # profile-permitted exception, independently of the photo integration tests.
    case.photo_exception = True
    sheets, profile = service.describe(REAL_TEMPLATE, legacy_profile())
    assert len(sheets) == 18
    assert profile == lermontov_profile()
    case.data.passport_number = "0000001"
    case.data.passport_issued_by = "MKK"
    case.data.passport_department_code = "000123"
    case.data.birth_date = date(1990, 1, 2)
    case.data.passport_issued_date = date(2020, 3, 4)
    case.data.living_address = Address(region="SYNTHETIC REGION", street="SYNTHETIC STREET 2")
    case.contract = ContractData(
        tools=[
            ToolItem(
                name=f"SYNTHETIC TOOL {i + 1}",
                inventory_number=f"000{i + 1}",
                quantity=i + 1,
                price_per_day=100 + i,
            )
            for i in range(tool_count)
        ],
        deposit_amount=0 if tool_count else None,
        purchase_price=999,
        actual_return_date=date(2026, 10, 17),
        actual_return_time="12:30",
    )
    case.confirmed_digest = case.digest()
    plan = service.plan(REAL_TEMPLATE, case, profile)
    target = service.write(plan, tmp_path / "result.xlsx", case, profile)
    assert file_hash(REAL_TEMPLATE) == before
    source, output = parts(REAL_TEMPLATE), parts(target)
    modified = {
        "xl/workbook.xml",
        "xl/_rels/workbook.xml.rels",
        "[Content_Types].xml",
        "xl/calcChain.xml",
        "xl/styles.xml",
    }
    editor_sheets = {
        "Дог.": "xl/worksheets/sheet1.xml",
        "Пр №1": "xl/worksheets/sheet2.xml",
        "Акт прием": "xl/worksheets/sheet3.xml",
        "Акт расч": "xl/worksheets/sheet4.xml",
    }
    modified.update(editor_sheets[c.sheet] for c in plan.changes if c.sheet in editor_sheets)
    for name in source.keys() - modified:
        assert output[name] == source[name], name
    # All nodes outside planned cells remain semantically identical, including
    # namespace declarations, print configuration, merges and formula caches.
    for sheet, path in editor_sheets.items():
        expected = {c.cell for c in plan.changes if c.sheet == sheet}
        if not expected:
            continue
        docs = [document(data[path]) for data in [source, output]]
        for doc in docs:
            for cell in list(doc.getElementsByTagNameNS(MAIN, "c")):
                if cell.getAttribute("r") in expected:
                    cell.parentNode.removeChild(cell)
        assert docs[0].documentElement.toxml() == docs[1].documentElement.toxml()
    book = load_workbook(target, keep_links=True)
    assert len(book._external_links) == 1
    assert book["Дог."]["K19"].value == "000123"
    assert book["Дог."]["G19"].value == "0000001"
    assert book["Акт прием"]["N21"].value == 999
    assert book["Акт расч"]["N20"].value == "='Пр №1'!N20"
    for i in range(tool_count):
        page, offset = divmod(i, 5)
        suffix = f" ({page + 1})" if page else ""
        app = book["Пр №1" + suffix]
        act = book["Акт прием" + suffix]
        assert app[f"B{20 + offset * 2}"].value == f"SYNTHETIC TOOL {i + 1}"
        assert act[f"L{21 + offset}"].value == f"000{i + 1}"
        assert app["J20"].value == "=NOW()"
        assert len(app._images) == len(act._images) == 1
        if page:
            assert act["B21"].value == f"='Пр №1{suffix}'!B20"
            assert book["Акт расч" + suffix]["N20"].value == f"='Пр №1{suffix}'!N20"
            assert app["P20"].value is None
            assert "_xlnm.Print_Area" not in act.defined_names
            assert act.print_area.endswith("$A$1:$Q$52")
    book.close()


def test_custom_profile_is_not_replaced(tmp_path, template):
    profile = legacy_profile()
    profile.name = "My customized profile"
    _, selected = WorkbookService().describe(template, profile)
    assert selected == profile
