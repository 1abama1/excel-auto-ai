"""Synthetic photographs only; actual template is read without personal output."""

from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image, ImageDraw

from document_excel.domain.errors import AppError
from document_excel.domain.models import ClientCase, DocumentSide, SourceImage
from document_excel.infrastructure.excel import WorkbookService, file_hash
from document_excel.infrastructure.excel_images import SheetGeometry
from document_excel.presentation.dialogs import ExportDialog


@pytest.fixture
def photo_setup(template, case, profile, tmp_path):
    book = load_workbook(template)
    sheet = book.create_sheet("Пасп Доп Данные")
    for column in "ABCDEFGHIJK":
        sheet.column_dimensions[column].width = 12 if column < "F" else 15
    sheet.row_dimensions[5].height = 31
    sheet.merge_cells("A1:E4")
    sheet["N32"] = "KEEP"
    sheet["N31"] = "=1+2"
    sheet.print_area = "A1:K24"
    book.save(template)
    book.close()
    for side, size, color in [
        (DocumentSide.FRONT, (900, 500), "lightblue"),
        (DocumentSide.BACK, (350, 850), "lightgreen"),
    ]:
        path = tmp_path / (side.value + ".png")
        image = Image.new("RGB", size, color)
        ImageDraw.Draw(image).text((20, 20), "SYNTHETIC " + side.value, fill="black")
        image.save(path)
        case.images.append(SourceImage(path=path, role=side))
    profile.photo_sheet = "пасп   доп данные"
    profile.photos_required = True
    case.data.document_type = "kg_id"
    case.confirmed_digest = case.digest()
    return template, case, profile


def image_bounds(geometry, anchor):
    return (
        *geometry.position(anchor._from.col, anchor._from.row, anchor._from.colOff, anchor._from.rowOff),
        *geometry.position(anchor.to.col, anchor.to.row, anchor.to.colOff, anchor.to.rowOff),
    )


def test_two_photos_fit_real_cell_geometry_and_preserve_workbook(photo_setup, tmp_path):
    template, case, profile = photo_setup
    before = file_hash(template)
    service = WorkbookService()
    plan = service.plan(template, case, profile)
    path = service.write(plan, tmp_path / "result.xlsx", case, profile)
    book = load_workbook(path)
    sheet = book["Пасп Доп Данные"]
    geometry = SheetGeometry(sheet)
    assert len(sheet._images) == 2
    boxes = []
    for image, expected in zip(sheet._images, plan.photos.insertions, strict=True):
        bounds = image_bounds(geometry, image.anchor)
        left, top, right, bottom = geometry.bounds(expected.area)
        assert left <= bounds[0] < bounds[2] <= right
        assert top <= bounds[1] < bounds[3] <= bottom
        original = Image.open(case.images[0 if expected.role == DocumentSide.FRONT else 1].path)
        assert (bounds[2] - bounds[0]) / (bounds[3] - bounds[1]) == pytest.approx(
            original.width / original.height, rel=1e-5
        )
        boxes.append(bounds)
    assert boxes[0][2] <= boxes[1][0]
    assert sheet["N31"].value == "=1+2" and sheet["N32"].value == "KEEP"
    assert sheet.print_area.endswith("$A$1:$K$24")
    assert "A1:E4" in {str(m) for m in sheet.merged_cells.ranges}
    assert file_hash(template) == before
    with ZipFile(path) as archive:
        photos = [n for n in archive.namelist() if n.startswith("xl/media/document-excel-")]
        assert len(photos) == 2 and all(archive.read(n).startswith(b"\x89PNG") for n in photos)
    book.close()


def test_repeat_export_replaces_owned_photos_without_duplicates(photo_setup, tmp_path):
    template, case, profile = photo_setup
    service = WorkbookService()
    first = service.write(service.plan(template, case, profile), tmp_path / "one.xlsx", case, profile)
    Image.new("RGB", (700, 400), "orange").save(case.images[1].path)
    case.confirmed_digest = case.digest()
    second = service.write(service.plan(first, case, profile), tmp_path / "two.xlsx", case, profile)
    book = load_workbook(second)
    assert len(book["Пасп Доп Данные"]._images) == 2
    with ZipFile(second) as archive:
        assert len([n for n in archive.namelist() if n.startswith("xl/media/document-excel-")]) == 2
    book.close()


def test_unknown_overlap_needs_explicit_consent_and_other_logo_survives(photo_setup, tmp_path, qtbot):
    template, case, profile = photo_setup
    book = load_workbook(template)
    for position in ["A3", "M30"]:
        logo = ExcelImage(case.images[0].path)
        logo.width, logo.height = 60, 40
        book["Пасп Доп Данные"].add_image(logo, position)
    book.save(template)
    book.close()
    service = WorkbookService()
    plan = service.plan(template, case, profile)
    assert len(plan.photos.conflicts) == 1
    with pytest.raises(AppError, match="Подтвердите"):
        service.write(plan, tmp_path / "rejected.xlsx", case, profile)
    assert not (tmp_path / "rejected.xlsx").exists()
    dialog = ExportDialog(plan)
    qtbot.addWidget(dialog)
    dialog.confirm.setChecked(True)
    assert not dialog.save_button.isEnabled()
    dialog.image_confirm.setChecked(True)
    assert dialog.save_button.isEnabled()
    dialog.accept()
    assert plan.image_replace_confirmed
    path = service.write(plan, tmp_path / "accepted.xlsx", case, profile)
    book = load_workbook(path)
    assert len(book["Пасп Доп Данные"]._images) == 3
    book.close()


@pytest.mark.parametrize(
    "reason", ["front_missing", "duplicate", "sheet_missing", "file_changed", "other_client"]
)
def test_unsafe_photo_export_is_rejected(photo_setup, tmp_path, reason):
    template, case, profile = photo_setup
    service = WorkbookService()
    if reason == "front_missing":
        case.images.pop(0)
    elif reason == "duplicate":
        case.images.append(case.images[0].model_copy())
    elif reason == "sheet_missing":
        profile.photo_sheet = "missing"
    elif reason == "other_client":
        template = service.write(
            service.plan(template, case, profile), tmp_path / "first.xlsx", case, profile
        )
        case.id = "another-client"
    case.confirmed_digest = case.digest()
    if reason == "file_changed":
        plan = service.plan(template, case, profile)
        Image.new("RGB", (900, 500), "red").save(case.images[0].path)
        with pytest.raises(AppError):
            service.write(plan, tmp_path / "changed.xlsx", case, profile)
    else:
        with pytest.raises(AppError):
            service.plan(template, case, profile)


def test_missing_photo_exception_is_case_specific(photo_setup, tmp_path):
    template, case, profile = photo_setup
    case.images = []
    case.photo_exception = True
    profile.allow_photo_exception = True
    case.confirmed_digest = case.digest()
    service = WorkbookService()
    plan = service.plan(template, case, profile)
    assert not plan.photos.insertions
    service.write(plan, tmp_path / "excepted.xlsx", case, profile)
    other = ClientCase(data=case.data, registration_status=case.registration_status)
    other.confirmed_digest = other.digest()
    with pytest.raises(AppError):
        service.plan(template, other, profile)


def test_confirmed_processed_photo_is_embedded_instead_of_original(photo_setup, tmp_path):
    import io

    import numpy as np

    from document_excel.application.workflows import CaseService
    from document_excel.infrastructure.document_images import DocumentCaptureController

    template, case, profile = photo_setup
    controller = DocumentCaptureController(CaseService())
    original = case.images[0]
    photo = controller.photo(DocumentSide.FRONT, np.full((300, 500, 3), 200, np.uint8), source=original)
    controller.assign(case, photo)
    case.confirmed_digest = case.digest()
    service = WorkbookService()
    plan = service.plan(template, case, profile)
    front = next(p for p in plan.photos.insertions if p.role == DocumentSide.FRONT)
    assert Image.open(io.BytesIO(front.data)).size == (500, 300)
    path = service.write(plan, tmp_path / "processed.xlsx", case, profile)
    controller.cleanup()
    assert original.path.exists()
    book = load_workbook(path)
    assert len(book["Пасп Доп Данные"]._images) == 2
    book.close()


def test_exception_removes_only_owned_stale_photo_and_unused_media(photo_setup, tmp_path):
    template, case, profile = photo_setup
    service = WorkbookService()
    first = service.write(service.plan(template, case, profile), tmp_path / "complete.xlsx", case, profile)
    case.images = [s for s in case.images if s.role == DocumentSide.FRONT]
    case.photo_exception = True
    profile.allow_photo_exception = True
    case.confirmed_digest = case.digest()
    plan = service.plan(first, case, profile)
    assert plan.photos.exception_used
    second = service.write(plan, tmp_path / "missing-back.xlsx", case, profile)
    book = load_workbook(second)
    assert len(book["Пасп Доп Данные"]._images) == 1
    book.close()
    with ZipFile(second) as archive:
        assert len([n for n in archive.namelist() if n.startswith("xl/media/document-excel-")]) == 1
