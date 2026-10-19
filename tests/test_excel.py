from datetime import date, datetime
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from document_excel.domain.errors import AppError
from document_excel.domain.models import Address, ContractData, RegistrationStatus, ToolItem
from document_excel.infrastructure.excel import (
    CellMapping,
    TemplateProfile,
    TemplateRegistry,
    WorkbookService,
    file_hash,
    legacy_profile,
    readiness,
)


def alter(path, callback):
    book = load_workbook(path)
    callback(book)
    book.save(path)
    book.close()


def test_roundtrip_preserves_original_styles_formulas_and_pin(template, case, profile, tmp_path):
    service = WorkbookService()
    original = file_hash(template)
    plan = service.plan(template, case, profile)
    output = service.write(plan, tmp_path / "result.xlsx", case, profile)
    assert file_hash(template) == original
    book = load_workbook(output)
    source = load_workbook(template)
    assert book["Дог."]["C22"].value == "00123456789012"
    assert book["Дог."]["C22"].data_type == "s"
    assert book["Дог."]["C22"]._style == source["Дог."]["C22"]._style
    assert book["Акт расч"]["L21"].value == "=G21-F21"
    assert book["Дог."].print_area == source["Дог."].print_area
    assert list(book["Дог."].merged_cells) == list(source["Дог."].merged_cells)
    book.close()
    source.close()


def test_formula_protection(template, case, profile):
    alter(template, lambda b: setattr(b["Дог."]["A18"], "value", "=1+1"))
    with pytest.raises(AppError, match="формулу"):
        WorkbookService().plan(template, case, profile)


def test_explicit_formula_replacement(template, case, profile, tmp_path):
    alter(template, lambda b: setattr(b["Дог."]["A18"], "value", "=1+1"))
    next(m for m in profile.field_mappings if m.field == "full_name").allow_formula_replace = True
    service = WorkbookService()
    plan = service.plan(template, case, profile)
    assert any("формулы" in c.warning for c in plan.changes)
    service.write(plan, tmp_path / "result.xlsx", case, profile)


def test_merged_nonleading_is_blocked(template, case, profile):
    alter(template, lambda b: b["Дог."].merge_cells("F19:G19"))
    with pytest.raises(AppError, match="неведущая"):
        WorkbookService().plan(template, case, profile)


def test_merged_leading_is_supported(template, case, profile, tmp_path):
    alter(template, lambda b: b["Дог."].merge_cells("G19:I19"))
    service = WorkbookService()
    service.write(service.plan(template, case, profile), tmp_path / "result.xlsx", case, profile)


def test_missing_sheet(template, case, profile):
    alter(template, lambda b: b.remove(b["Дог."]))
    with pytest.raises(AppError, match="обязательный лист"):
        WorkbookService().plan(template, case, profile)


def test_formula_injection_written_as_text(template, case, profile, tmp_path):
    case.data.full_name = '=HYPERLINK("https://invalid", "x")'
    case.confirmed_digest = case.digest()
    service = WorkbookService()
    path = service.write(service.plan(template, case, profile), tmp_path / "result.xlsx", case, profile)
    book = load_workbook(path)
    assert book["Дог."]["A18"].data_type == "s"
    assert book["Дог."]["A18"].value.startswith("=")
    book.close()


def test_more_than_five_tools(template, case, profile, tmp_path):
    alter(template, lambda b: setattr(b["Пр №1"], "print_title_rows", "1:3"))
    case.contract = ContractData(
        tools=[ToolItem(name=f"Tool {i}", inventory_number=f"000{i}") for i in range(7)], deposit_amount=150
    )
    case.confirmed_digest = case.digest()
    service = WorkbookService()
    plan = service.plan(template, case, profile)
    path = service.write(plan, tmp_path / "tools.xlsx", case, profile)
    book = load_workbook(path)
    assert book["Пр №1"]["P20"].value == 150
    assert book["Пр №1 (2)"]["P20"].value is None
    assert book["Пр №1 (2)"]["B20"].value == "Tool 5"
    assert book["Пр №1 (2)"]["K20"].value == "0005"
    assert book["Пр №1 (2)"]["B22"].value is None
    assert book["Пр №1 (2)"].print_area.endswith("$A$1:$P$30")
    assert book["Пр №1 (2)"].print_title_rows == "$1:$3"
    book.close()


@pytest.mark.parametrize(
    "reason", ["same_path", "exists", "changed_data", "changed_source", "changed_profile"]
)
def test_stale_or_unsafe_save_is_rejected(reason, template, case, profile, tmp_path):
    service = WorkbookService()
    plan = service.plan(template, case, profile)
    destination = tmp_path / "output.xlsx"
    if reason == "same_path":
        destination = template
    elif reason == "exists":
        destination.write_bytes(b"keep me")
    elif reason == "changed_data":
        case.data.full_name = "CHANGED"
    elif reason == "changed_source":
        alter(template, lambda b: setattr(b["Дог."]["P30"], "value", "changed"))
    else:
        profile.phone_mode = "separate"
    with pytest.raises(AppError):
        service.write(plan, destination, case, profile)
    if reason == "exists":
        assert destination.read_bytes() == b"keep me"


def test_atomic_publish_failure_cleans_temp(template, case, profile, tmp_path, monkeypatch):
    service = WorkbookService()
    plan = service.plan(template, case, profile)

    def fail(*args):
        raise PermissionError("blocked")

    monkeypatch.setattr("document_excel.infrastructure.excel.os.link", fail)
    target = tmp_path / "failure.xlsx"
    with pytest.raises(AppError):
        service.write(plan, target, case, profile)
    assert not target.exists()
    assert not list(tmp_path.glob(".document-excel-*"))


def test_blank_manual_fields_are_preserved(template, case, profile, tmp_path):
    alter(template, lambda b: setattr(b["Дог."]["A21"], "value", "MANUAL LIVING ADDRESS"))
    service = WorkbookService()
    output = service.write(service.plan(template, case, profile), tmp_path / "result.xlsx", case, profile)
    book = load_workbook(output)
    assert book["Дог."]["A21"].value == "MANUAL LIVING ADDRESS"
    book.close()


def test_phone_variants_do_not_duplicate(template, case):
    case.data.whatsapp_phone = "+996 555 000001"
    case.data.additional_phone = "+996 555 000002"
    case.confirmed_digest = case.digest()
    service = WorkbookService()
    combined = service.plan(template, case, legacy_profile("combined"))
    assert [c.cell for c in combined.changes if c.cell in {"D7", "G7"}] == ["D7"]
    separated = service.plan(template, case, legacy_profile("separate"))
    assert {c.cell for c in separated.changes if c.cell in {"D7", "G7"}} == {"D7", "G7"}


def test_profile_validation_and_roundtrip(tmp_path, profile):
    registry = TemplateRegistry(tmp_path)
    assert registry.load(registry.save(profile)) == profile
    with pytest.raises(ValidationError):
        CellMapping(field="full_name", sheet="Дог.", cell="XFE1")
    with pytest.raises(ValidationError):
        TemplateProfile(
            field_mappings=[
                CellMapping(field="phones", sheet="A", cell="A1"),
                CellMapping(field="full_name", sheet="A", cell="A1"),
            ]
        )


def test_missing_registration_requires_business_exception(case, profile):
    case.data.registration_address = Address()
    case.registration_status = RegistrationStatus.NOT_PRESENT
    case.confirmed_digest = case.digest()
    assert readiness(case, profile)
    profile.allow_missing_registration = True
    assert not readiness(case, profile)


def test_date_type_and_return_formula(template, case, profile, tmp_path):
    case.contract.actual_return_date = date(2026, 1, 10)
    case.confirmed_digest = case.digest()
    profile.enable_return_formula = True
    service = WorkbookService()
    result = service.write(service.plan(template, case, profile), tmp_path / "return.xlsx", case, profile)
    book = load_workbook(result)
    assert book["Акт расч"]["G21"].value == datetime(2026, 1, 10)
    assert book["Акт расч"]["L21"].value == "=G21-F21"
    book.close()


def test_special_excel_objects_rejected(template, case, profile):
    with ZipFile(template, "a", ZIP_DEFLATED) as archive:
        archive.writestr("xl/activeX/fake.xml", "fake")
    with pytest.raises(AppError, match="специальные объекты"):
        WorkbookService().plan(template, case, profile)


def test_custom_mapping(template, case, tmp_path):
    profile = TemplateProfile(
        required_sheets=["Дог."],
        tool_sheet=None,
        field_mappings=[CellMapping(field="personal_identifier", sheet="Дог.", cell="B29")],
    )
    service = WorkbookService()
    target = service.write(service.plan(template, case, profile), tmp_path / "custom.xlsx", case, profile)
    book = load_workbook(target)
    assert book["Дог."]["B29"].value == "00123456789012"
    book.close()
