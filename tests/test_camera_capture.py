import hashlib
import threading
import time

import cv2
import numpy as np
import pytest
from PIL import Image
from PySide6.QtCore import QObject, Signal

from document_excel.application.workflows import CaseService
from document_excel.domain.models import ClientCase, DocumentSide, SourceImage
from document_excel.infrastructure.camera import CameraDevice, CameraService
from document_excel.infrastructure.document_images import (
    DocumentCaptureController,
    DocumentImageProcessor,
    selected_image,
)
from document_excel.presentation.camera_dialog import CameraDialog, PhotoEditor


class MockCapture:
    def __init__(self, opened=True):
        self.opened = opened
        self.width, self.height = 640, 480
        self.released = False
        self.fail = False
        self.threads = []

    def isOpened(self):
        return self.opened

    def set(self, prop, value):
        self.threads.append(threading.get_ident())
        if prop == cv2.CAP_PROP_FRAME_WIDTH:
            self.width = min(1920, value)
        if prop == cv2.CAP_PROP_FRAME_HEIGHT:
            self.height = min(1080, value)
        return True

    def read(self):
        self.threads.append(threading.get_ident())
        time.sleep(0.002)
        return (False, None) if self.fail else (True, np.full((self.height, self.width, 3), 210, np.uint8))

    def release(self):
        self.released = True


def test_camera_discovery_is_mockable_and_releases_probes(qtbot):
    captures = []

    def factory(index, backend):
        cap = MockCapture(index in {0, 2})
        captures.append(cap)
        return cap

    service = CameraService(capture_factory=factory, indices=range(3))
    with qtbot.waitSignal(service.devices_found) as signal:
        service.discover()
    qtbot.waitUntil(lambda: not service.busy)
    assert [d.index for d in signal.args[0]] == [0, 2]
    assert all(c.released for c in captures)


def test_capture_uses_full_resolution_and_background_thread(qtbot):
    cap = MockCapture()
    service = CameraService(capture_factory=lambda *_: cap)
    previews = []
    service.preview.connect(previews.append)
    try:
        with qtbot.waitSignal(service.ready, timeout=5000):
            service.start(CameraDevice(0, "Synthetic camera"))
        with qtbot.waitSignal(service.photograph, timeout=5000) as photograph:
            service.capture()
        qtbot.waitUntil(lambda: bool(previews))
        assert photograph.args[0].shape[:2] == (1080, 1920)
        assert max(previews[-1].shape[:2]) <= 1000
        assert all(t != threading.get_ident() for t in cap.threads)
    finally:
        service.close()
        qtbot.waitUntil(lambda: not service.busy, timeout=5000)
    assert cap.released


def test_camera_disconnect_reports_error_and_releases_device(qtbot):
    cap = MockCapture()
    service = CameraService(capture_factory=lambda *_: cap)
    with qtbot.waitSignal(service.ready, timeout=5000):
        service.start(CameraDevice(0, "Synthetic camera"))
    with qtbot.waitSignal(service.error, timeout=5000) as error:
        cap.fail = True
    qtbot.waitUntil(lambda: not service.busy)
    assert "прерван" in error.args[0]
    assert cap.released


def test_camera_unavailable_or_permission_denied_is_actionable(qtbot):
    service = CameraService(capture_factory=lambda *_: MockCapture(False))
    with qtbot.waitSignal(service.error) as error:
        service.start(CameraDevice(0, "Unavailable"))
    qtbot.waitUntil(lambda: not service.busy)
    assert "доступ" in error.args[0] and "камер" in error.args[0].lower()


class FakeCameraService(QObject):
    devices_found = Signal(object)
    preview = Signal(object)
    photograph = Signal(object)
    ready = Signal(object)
    error = Signal(str)
    busy = False
    closed = False

    def discover(self):
        self.devices_found.emit([CameraDevice(0, "Synthetic camera")])

    def start(self, *_):
        self.ready.emit([(900, 600)])
        self.preview.emit(np.full((300, 450, 3), 200, np.uint8))

    def capture(self):
        self.photograph.emit(np.full((600, 900, 3), 200, np.uint8))

    def stop(self):
        pass

    def close(self):
        self.closed = True


def test_physical_camera_mode_emphasizes_mrz_and_keeps_retake_unaccepted(qtbot):
    service = FakeCameraService()
    dialog = CameraDialog(service=service, physical=True, side=DocumentSide.BACK)
    qtbot.addWidget(dialog)
    dialog.show()
    qtbot.waitUntil(lambda: dialog.shutter.isEnabled())
    assert "MRZ" in dialog.hint.text() and "бликов" in dialog.hint.text()
    assert dialog.preview_quality.text()
    accepted = []
    dialog.photo_accepted.connect(lambda *args: accepted.append(args))
    dialog.capture()
    dialog.resume()
    dialog.reject()
    qtbot.waitUntil(lambda: not dialog.isVisible())
    assert not accepted and service.closed


def test_camera_dialog_front_back_retake_and_close(qtbot):
    service = FakeCameraService()
    dialog = CameraDialog(service=service)
    qtbot.addWidget(dialog)
    accepted = []
    dialog.photo_accepted.connect(
        lambda side, raw, processed: accepted.append((side, raw.shape, processed.shape))
    )
    dialog.show()
    qtbot.waitUntil(lambda: dialog.shutter.isEnabled())
    dialog.capture()
    assert dialog.stack.currentIndex() == 1
    dialog.resume()
    assert not accepted
    dialog.capture()
    dialog.use_photo()
    assert accepted[0][0] == DocumentSide.FRONT
    assert dialog.side.currentData() == DocumentSide.BACK
    dialog.capture()
    dialog.use_photo()
    qtbot.waitUntil(lambda: not dialog.isVisible())
    assert [row[0] for row in accepted] == [DocumentSide.FRONT, DocumentSide.BACK]
    assert service.closed


def test_cancel_camera_does_not_accept_an_unconfirmed_frame(qtbot):
    service = FakeCameraService()
    dialog = CameraDialog(service=service)
    qtbot.addWidget(dialog)
    accepted = []
    dialog.photo_accepted.connect(lambda *_: accepted.append(True))
    dialog.show()
    qtbot.waitUntil(lambda: dialog.shutter.isEnabled())
    dialog.capture()
    dialog.reject()
    qtbot.waitUntil(lambda: not dialog.isVisible())
    assert not accepted and service.closed


def test_confirmed_photos_are_owned_by_case_and_retake_keeps_other_side(tmp_path):
    controller = DocumentCaptureController(CaseService())
    first, other = ClientCase(), ClientCase()
    raw = np.full((800, 1100, 3), 220, np.uint8)
    front = controller.photo(DocumentSide.FRONT, raw[100:700, 100:1000], original=raw)
    back = controller.photo(DocumentSide.BACK, raw, original=raw)
    controller.assign(first, front)
    controller.assign(first, back)
    original_hash = hashlib.sha256(front.path.read_bytes()).hexdigest()
    replacement = controller.photo(DocumentSide.BACK, raw[:700], original=raw)
    controller.assign(first, replacement)
    assert len(first.images) == 2 and first.images[0].id == front.id
    assert first.images[1].id == replacement.id and back.id not in [s.id for s in first.images]
    assert not other.images
    assert front.path != front.processed_path
    assert selected_image(front).shape == (600, 900, 3)
    assert hashlib.sha256(front.path.read_bytes()).hexdigest() == original_hash
    controller.swap(first)
    assert front.role == DocumentSide.BACK and replacement.role == DocumentSide.FRONT
    controller.cleanup()
    assert not front.path.exists()


def test_mixed_disk_and_camera_images_persist_original_and_processed(tmp_path):
    controller = DocumentCaptureController(CaseService())
    original = tmp_path / "file.png"
    Image.new("RGB", (1000, 700), "white").save(original)
    before = original.read_bytes()
    source = SourceImage(path=original, role=DocumentSide.BACK)
    edited = controller.photo(
        DocumentSide.BACK,
        np.full((500, 900, 3), 200, np.uint8),
        source=source,
        retain=tmp_path / "saved",
        case_id="case-a",
    )
    controller.cleanup()
    assert edited.path.exists() and edited.processed_path.exists()
    assert original.read_bytes() == edited.path.read_bytes() == before
    assert edited.processed_path.read_bytes() != before


def test_photo_editor_manual_crop_rotation_and_reset(qtbot):
    editor = PhotoEditor()
    qtbot.addWidget(editor)
    image = np.full((800, 1100, 3), 230, np.uint8)
    editor.load(image)
    editor.crop((100, 200, 700, 400))
    assert editor.current.shape[:2] == (400, 700)
    editor.rotate()
    assert editor.current.shape[:2] == (700, 400)
    editor.reset()
    assert np.array_equal(editor.current, image)


def test_dark_blurred_photo_quality_warning():
    notes = DocumentImageProcessor.quality(np.full((800, 1200, 3), 20, np.uint8))
    assert any("размытие" in n for n in notes)
    assert any("света" in n for n in notes)


def test_auto_perspective_suggestion_keeps_original_pixels():
    original = np.zeros((800, 1200, 3), np.uint8)
    cv2.fillConvexPoly(
        original, np.array([[100, 150], [1100, 100], [1050, 700], [150, 720]]), (240, 240, 240)
    )
    before = original.copy()
    processed = DocumentImageProcessor().suggest(original)
    assert np.array_equal(original, before)
    assert processed.shape != original.shape
    assert min(processed.shape[:2]) >= 500


def test_main_window_waits_for_camera_before_cleaning_capture_files(qtbot, tmp_path, monkeypatch):
    from PySide6.QtCore import QTimer

    from document_excel.infrastructure.storage import SettingsRepository
    from document_excel.presentation.main_window import MainWindow

    monkeypatch.setattr(MainWindow, "diagnose", lambda self: None)
    window = MainWindow(SettingsRepository(tmp_path / "close-app"))
    qtbot.addWidget(window)

    class DelayedCamera(FakeCameraService):
        def close(self):
            self.closed = True
            QTimer.singleShot(80, lambda: setattr(self, "busy", False))

    service = DelayedCamera()
    dialog = CameraDialog(window, service=service)
    raw = np.full((700, 1000, 3), 210, np.uint8)
    photo = window.capture_controller.photo(DocumentSide.FRONT, raw, original=raw)
    window.show()
    dialog.show()
    qtbot.waitUntil(lambda: dialog.shutter.isEnabled())
    service.busy = True
    window.close()
    assert photo.path.exists()
    assert window.isVisible()
    qtbot.waitUntil(lambda: not window.isVisible(), timeout=3000)
    assert service.closed and not service.busy
    assert not photo.path.exists()


@pytest.mark.ocr
def test_camera_ui_to_real_ocr_and_embedded_excel(qtbot, image_factory, template, tmp_path, monkeypatch):
    from openpyxl import load_workbook
    from PySide6.QtCore import QTimer

    from document_excel.domain.models import RegistrationStatus
    from document_excel.infrastructure.ocr import TesseractEngine
    from document_excel.infrastructure.storage import Settings, SettingsRepository
    from document_excel.presentation import camera_dialog
    from document_excel.presentation.main_window import MainWindow

    settings = Settings(ocr_engine="tesseract", languages=["eng"], preprocessing="original")
    if TesseractEngine(settings).diagnose().error:
        pytest.skip("Tesseract is unavailable")
    array = np.asarray(Image.open(image_factory()).convert("RGB"))

    class SyntheticCamera(FakeCameraService):
        def capture(self):
            self.photograph.emit(array.copy())

    monkeypatch.setattr(MainWindow, "diagnose", lambda self: None)
    window = MainWindow(SettingsRepository(tmp_path / "capture-app"))
    window.settings.ocr_engine = "tesseract"
    window.settings.languages = ["eng"]
    window.settings.preprocessing = "original"
    book = load_workbook(template)
    book.create_sheet("Пасп Доп Данные")
    book.save(template)
    book.close()
    window.current_case.template_path = template
    _, window.profile = window.workbooks.describe(template, window.profile, strict=True)
    window.current_case.workbook_profile = window.profile.model_dump(mode="json")
    errors = []
    monkeypatch.setattr(window, "error", errors.append)
    qtbot.addWidget(
        window,
        before_close_func=lambda _: qtbot.waitUntil(lambda: not window.background_workers, timeout=20000),
    )
    service = SyntheticCamera()
    dialog = CameraDialog(window, service=service)
    monkeypatch.setattr(camera_dialog, "CameraDialog", lambda *_args, **_kwargs: dialog)

    def take_both():
        if not dialog.shutter.isEnabled():
            QTimer.singleShot(25, take_both)
            return
        for _ in range(2):
            dialog.capture()
            dialog.use_photo()

    QTimer.singleShot(150, take_both)
    window.case_actions.capture_documents()
    assert window.current_case.data.personal_identifier is None
    window.case_actions.start_ocr()  # Capture returns to upload; OCR is an explicit action.
    qtbot.waitUntil(
        lambda: bool(window.current_case.data.personal_identifier) and not window.background_workers,
        timeout=20000,
    )
    case = window.current_case
    assert not errors and service.closed
    assert {s.role for s in case.images} == {DocumentSide.FRONT, DocumentSide.BACK}
    assert all(s.processed_confirmed for s in case.images)
    assert case.data.personal_identifier == "00123456789012"
    assert case.data.passport_number == "0000001"
    window.apply_field_edit("full_name", "Тестов Пример Тестович")
    case.registration_status = RegistrationStatus.USER_CONFIRMED
    case.confirmed_digest = case.digest()
    result = window.workbooks.write(
        window.workbooks.plan(template, case, window.profile),
        tmp_path / "capture-result.xlsx",
        case,
        window.profile,
    )
    book = load_workbook(result)
    assert len(book["Пасп Доп Данные"]._images) == 2
    assert book["Дог."]["C22"].value == "00123456789012"
    book.close()
