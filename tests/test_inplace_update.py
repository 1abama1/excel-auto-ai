"""The chosen XLSX is replaced only after a complete verified write."""

import errno

import pytest
from openpyxl import load_workbook
from PIL import Image
from test_best_effort import old_client
from test_best_effort import setup_copy as setup_copy

from document_excel.domain.errors import AppError
from document_excel.domain.models import DocumentSide, MissingValuePolicy, SourceImage
from document_excel.infrastructure.excel import WorkbookService, file_hash


def prepared(setup_copy):
    source, case, profile = setup_copy
    case.data.full_name = "ИВАНОВ ИВАН"
    service = WorkbookService()
    plan = service.plan(source, case, profile)
    plan.select_missing_policy(MissingValuePolicy.PRESERVE_EXISTING)
    return service, plan, source, case, profile


def test_update_selected_file_keeps_name_and_other_data(setup_copy, tmp_path):
    service, plan, source, case, profile = prepared(setup_copy)
    before = file_hash(source)
    output = service.update_existing(plan, case, profile)
    assert output == source and file_hash(source) != before
    assert list(tmp_path.glob("*.xlsx")) == [source]
    assert not list(tmp_path.glob(".document-excel-*"))
    book = load_workbook(source)
    assert book["Дог."]["A18"].value == "Иванов Иван"
    assert book["Пр №1"]["N20"].value == 350
    assert book["Пр №1"]["P20"].value == 5000
    assert book["Акт расч"]["L21"].value == "=G21-F21"
    book.close()
    case.data.full_name = "ИВАНОВ ПЁТР"
    next_plan = service.plan(source, case, profile)
    next_plan.select_missing_policy(MissingValuePolicy.CLEAR_OLD_PERSONAL)
    assert service.update_existing(next_plan, case, profile) == source
    assert list(tmp_path.glob("*.xlsx")) == [source]


@pytest.mark.parametrize("failure", ["save", "verify", "replace"])
def test_failed_update_keeps_original_and_cleans_temp(setup_copy, tmp_path, monkeypatch, failure):
    from document_excel.infrastructure.excel_images import ExcelImageEmbedder
    from document_excel.infrastructure.xlsx_package import PackageEditor

    service, plan, source, case, profile = prepared(setup_copy)
    before = source.read_bytes()
    if failure == "save":

        def fail(*args):
            raise OSError(errno.ENOSPC, "full disk")

        monkeypatch.setattr(PackageEditor, "save", fail)
    elif failure == "verify":

        def fail(*args):
            raise AppError("Результат не прошёл проверку")

        monkeypatch.setattr(ExcelImageEmbedder, "verify", fail)
    else:

        def fail(*args):
            raise PermissionError(errno.EACCES, "workbook open in Excel")

        monkeypatch.setattr("document_excel.infrastructure.excel.os.replace", fail)
    with pytest.raises(AppError):
        service.update_existing(plan, case, profile)
    assert source.read_bytes() == before
    assert list(tmp_path.glob("*.xlsx")) == [source]
    assert not list(tmp_path.glob(".document-excel-*"))


def test_external_change_immediately_before_commit_is_preserved(setup_copy, tmp_path, monkeypatch):
    import shutil

    service, plan, source, case, profile = prepared(setup_copy)
    changed = source.read_bytes() + b"external change"
    original = shutil.copymode

    def changed_during_commit(*args):
        original(*args)
        source.write_bytes(changed)

    monkeypatch.setattr("document_excel.infrastructure.excel.shutil.copymode", changed_during_commit)
    with pytest.raises(AppError, match="изменился перед записью"):
        service.update_existing(plan, case, profile)
    assert source.read_bytes() == changed
    assert not list(tmp_path.glob(".document-excel-*"))


def test_modified_input_and_invalid_result_never_replace_selected_file(setup_copy, tmp_path):
    service, plan, source, case, profile = prepared(setup_copy)
    source.write_bytes(source.read_bytes() + b"external change")
    before = source.read_bytes()
    with pytest.raises(AppError, match="изменился"):
        service.update_existing(plan, case, profile)
    assert source.read_bytes() == before


def test_update_cannot_target_another_selected_document(setup_copy, tmp_path):
    service, plan, source, case, profile = prepared(setup_copy)
    before = source.read_bytes()
    case.template_path = tmp_path / "other.xlsx"
    with pytest.raises(AppError, match="тот же"):
        service.update_existing(plan, case, profile)
    assert source.read_bytes() == before
    case.template_path = source
    with pytest.raises(AppError, match="только выбранный"):
        service.write(plan, tmp_path / "other.xlsx", case, profile, replace_source=True)


def test_new_client_partial_photo_updates_original_and_removes_old_side(setup_copy, tmp_path):
    source, case, profile = old_client(setup_copy, tmp_path)
    case.data.full_name = "НОВЫЙ КЛИЕНТ"
    front = tmp_path / "new-front.png"
    Image.new("RGB", (800, 500), "blue").save(front)
    case.images = [SourceImage(path=front, role=DocumentSide.FRONT)]
    service = WorkbookService()
    plan = service.plan(source, case, profile)
    plan.select_missing_policy(MissingValuePolicy.CLEAR_OLD_PERSONAL)
    plan.image_replace_confirmed = True
    assert service.update_existing(plan, case, profile) == source
    book = load_workbook(source)
    assert book["Дог."]["A18"].value == "Новый Клиент"
    assert book["Дог."]["C22"].value is None
    assert len(book[profile.photo_sheet]._images) == 2  # new front and unrelated logo
    book.close()
    assert list(tmp_path.glob("*.xlsx")) == [source]


def test_readonly_original_is_not_replaced(setup_copy, tmp_path):
    import stat

    service, plan, source, case, profile = prepared(setup_copy)
    before = source.read_bytes()
    mode = source.stat().st_mode
    try:
        source.chmod(stat.S_IREAD)
        with pytest.raises(AppError, match="только для чтения"):
            service.update_existing(plan, case, profile)
        assert source.read_bytes() == before
        assert not list(tmp_path.glob(".document-excel-*"))
    finally:
        source.chmod(mode)


def test_unknown_old_values_still_require_policy_for_inplace(setup_copy, tmp_path):
    source, case, profile = old_client(setup_copy, tmp_path)
    before = source.read_bytes()
    service = WorkbookService()
    plan = service.plan(source, case, profile)
    with pytest.raises(AppError, match="итоговой сводке"):
        service.update_existing(plan, case, profile)
    assert source.read_bytes() == before


def test_update_and_rename_moves_the_same_document(setup_copy, tmp_path):
    service, plan, source, case, profile = prepared(setup_copy)
    output = service.update_existing(plan, case, profile, new_name="Иванов Иван 20.10.2026.xlsx")
    assert output == tmp_path / "Иванов Иван 20.10.2026.xlsx"
    assert not source.exists() and list(tmp_path.glob("*.xlsx")) == [output]
    book = load_workbook(output)
    assert book["Дог."]["A18"].value == "Иванов Иван"
    assert book["Пр №1"]["P20"].value == 5000
    book.close()
    case.template_path = output
    repeated = service.plan(output, case, profile)
    repeated.select_missing_policy(MissingValuePolicy.PRESERVE_EXISTING)
    assert service.update_existing(repeated, case, profile, new_name=output.name) == output
    assert list(tmp_path.glob("*.xlsx")) == [output]


def test_rename_collision_keeps_existing_file_and_does_not_number_itself(setup_copy, tmp_path):
    service, plan, source, case, profile = prepared(setup_copy)
    occupied = tmp_path / "Иванов Иван 20.10.2026.xlsx"
    occupied.write_bytes(b"other client")
    output = service.update_existing(plan, case, profile, new_name=occupied.name)
    assert output.name == "Иванов Иван 20.10.2026 (2).xlsx"
    assert occupied.read_bytes() == b"other client" and not source.exists()
    case.template_path = output
    repeated = service.plan(output, case, profile)
    repeated.select_missing_policy(MissingValuePolicy.PRESERVE_EXISTING)
    assert service.update_existing(repeated, case, profile, new_name=occupied.name) == output
    assert len(list(tmp_path.glob("*.xlsx"))) == 2


def test_rename_is_rolled_back_if_atomic_write_fails(setup_copy, tmp_path, monkeypatch):
    service, plan, source, case, profile = prepared(setup_copy)
    before = source.read_bytes()

    def fail(*args):
        raise PermissionError(errno.EACCES, "blocked result")

    monkeypatch.setattr("document_excel.infrastructure.excel.os.replace", fail)
    with pytest.raises(AppError, match="обновить"):
        service.update_existing(plan, case, profile, new_name="Иванов Иван 20.10.2026.xlsx")
    assert source.read_bytes() == before
    assert list(tmp_path.glob("*.xlsx")) == [source]
    assert not list(tmp_path.glob(".document-excel-*"))


def test_failed_rename_keeps_original_content_and_name(setup_copy, tmp_path, monkeypatch):
    service, plan, source, case, profile = prepared(setup_copy)
    before = source.read_bytes()

    def fail(*args):
        raise PermissionError(errno.EACCES, "blocked rename")

    monkeypatch.setattr("document_excel.infrastructure.excel.rename_without_replace", fail)
    with pytest.raises(AppError, match="обновить"):
        service.update_existing(plan, case, profile, new_name="Иванов Иван 20.10.2026.xlsx")
    assert source.read_bytes() == before and list(tmp_path.glob("*.xlsx")) == [source]


def test_rename_race_uses_next_name_without_overwriting(setup_copy, tmp_path, monkeypatch):
    from document_excel.infrastructure import excel

    service, plan, source, case, profile = prepared(setup_copy)
    rename = excel.rename_without_replace
    occupied = tmp_path / "Иванов Иван 20.10.2026.xlsx"
    first = [True]

    def race(old, new):
        if first[0]:
            first[0] = False
            new.write_bytes(b"concurrent other file")
        return rename(old, new)

    monkeypatch.setattr(excel, "rename_without_replace", race)
    output = service.update_existing(plan, case, profile, new_name=occupied.name)
    assert output.name == "Иванов Иван 20.10.2026 (2).xlsx"
    assert occupied.read_bytes() == b"concurrent other file" and not source.exists()


@pytest.mark.parametrize("name", ["../other.xlsx", "Z:/other.xlsx", "CON.xlsx", "bad?.xlsx", "wrong.xls"])
def test_rename_accepts_only_safe_xlsx_name_in_same_folder(setup_copy, name):
    service, plan, source, case, profile = prepared(setup_copy)
    before = source.read_bytes()
    with pytest.raises(AppError, match="безопасное имя"):
        service.update_existing(plan, case, profile, new_name=name)
    assert source.read_bytes() == before


def test_failed_rollback_keeps_recoverable_file_at_new_path(setup_copy, tmp_path, monkeypatch):
    from document_excel.infrastructure import excel

    service, plan, source, case, profile = prepared(setup_copy)
    before = source.read_bytes()
    rename = excel.rename_without_replace
    calls = [0]

    def rename_then_block(old, new):
        calls[0] += 1
        if calls[0] > 1:
            raise PermissionError("rollback blocked")
        rename(old, new)

    def blocked_replace(*args):
        raise PermissionError("write blocked")

    monkeypatch.setattr(excel, "rename_without_replace", rename_then_block)
    monkeypatch.setattr(excel.os, "replace", blocked_replace)
    target = tmp_path / "Иванов Иван 20.10.2026.xlsx"
    with pytest.raises(AppError, match="Файл доступен по адресу"):
        service.update_existing(plan, case, profile, new_name=target.name)
    assert target.read_bytes() == before and not source.exists()
    assert list(tmp_path.glob("*.xlsx")) == [target]
    assert not list(tmp_path.glob(".document-excel-*"))
