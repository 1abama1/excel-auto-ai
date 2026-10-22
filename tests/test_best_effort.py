"""End-to-end safe copies with incomplete, synthetic identity data."""

import errno
from datetime import date
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image

from document_excel.domain.errors import AppError, OutputExistsError
from document_excel.domain.models import (
    Address,
    ClientCase,
    ClientData,
    DocumentSide,
    FieldEvidence,
    FieldStatus,
    MissingValuePolicy,
    SourceImage,
    WorkflowMode,
)
from document_excel.infrastructure.contract_filename import ContractFilenameGenerator
from document_excel.infrastructure.excel import (
    CellMapping,
    WorkbookService,
    create_demo_template,
    file_hash,
    legacy_profile,
    readiness,
    recognition_warnings,
)


@pytest.fixture
def setup_copy(tmp_path):
    path = tmp_path / "input.xlsx"
    create_demo_template(path)
    book = load_workbook(path)
    book["Пр №1"]["B20"] = "Перфоратор"
    book["Пр №1"]["N20"] = 350
    book["Пр №1"]["P20"] = 5000
    book.create_sheet("Пасп Доп Данные")["M30"] = "Логотип"
    book.save(path)
    book.close()
    service = WorkbookService()
    _, profile = service.describe(path, legacy_profile(), strict=True)
    case = ClientCase(
        workflow_mode=WorkflowMode.BEST_EFFORT_UPDATE,
        template_path=path,
        workbook_profile=profile.model_dump(mode="json"),
    )
    return path, case, profile


def old_client(setup_copy, tmp_path):
    path, case, profile = setup_copy
    book = load_workbook(path)
    for cell, value in {
        "A18": "Прежний Клиент",
        "C22": "00123456789012",
        "G19": "0000001",
        "J20": "Прежняя улица",
        "A21": "=C20",
        "I21": "=J20",
    }.items():
        book["Дог."][cell] = value
    for index, anchor in enumerate(["A1", "F1", "M1"]):
        photo = tmp_path / f"old-{index}.png"
        Image.new("RGB", (100, 60), "red").save(photo)
        book[profile.photo_sheet].add_image(ExcelImage(photo), anchor)
    book.save(path)
    book.close()
    return path, case, profile


def save(setup_copy, tmp_path, policy=MissingValuePolicy.PRESERVE_EXISTING):
    source, case, profile = setup_copy
    service = WorkbookService()
    plan = service.plan(source, case, profile)
    plan.select_missing_policy(policy)
    plan.image_replace_confirmed = True
    output = service.write(plan, tmp_path / "result.xlsx", case, profile)
    return plan, output


@pytest.mark.parametrize(
    "fields",
    [
        {"full_name": "ИВАНОВ ИВАН ИВАНОВИЧ"},
        {"passport_number": "0000002"},
        {"full_name": "ИВАНОВ ИВАН", "passport_number": "0000002"},
        {},
    ],
)
def test_partial_and_empty_data_are_savable(setup_copy, tmp_path, fields):
    source, case, profile = setup_copy
    case.data = ClientData(**fields)
    digest = file_hash(source)
    assert not readiness(case, profile)  # no OCR/confirmation/registration requirement
    warnings = recognition_warnings(case, profile)
    assert "registration_address" in warnings.missing_fields
    assert "personal_identifier" in warnings.missing_fields
    plan, output = save(setup_copy, tmp_path)
    assert file_hash(source) == digest
    book = load_workbook(output)
    assert book["Дог."]["A18"].value == (case.formatted_name or None)
    assert book["Дог."]["C22"].value is None
    assert book["Дог."]["J20"].value is None
    assert book["Пр №1"]["N20"].value == 350 and book["Пр №1"]["P20"].value == 5000
    assert book["Акт расч"]["L21"].value == "=G21-F21"
    book.close()
    if not fields:
        assert not plan.changes and file_hash(output) == digest


def test_full_document_updates_fields_and_embeds_photos(setup_copy, tmp_path):
    source, case, profile = setup_copy
    case.data = ClientData(
        full_name="ИВАНОВ ИВАН ИВАНОВИЧ",
        passport_number="0000002",
        personal_identifier="00987654321098",
        birth_date=date(1990, 1, 2),
        registration_address=Address(region="Тестовая область", street="Тестовая улица"),
    )
    for side in (DocumentSide.FRONT, DocumentSide.BACK):
        path = tmp_path / f"{side}.png"
        Image.new("RGB", (800, 500), "blue").save(path)
        case.images.append(SourceImage(path=path, role=side))
    plan, output = save(setup_copy, tmp_path)
    assert len(plan.photos.insertions) == 2
    book = load_workbook(output)
    assert len(book[profile.photo_sheet]._images) == 2
    assert book["Дог."]["C22"].value == "00987654321098"
    book.close()
    assert (
        ContractFilenameGenerator(lambda: date(2026, 10, 20)).generate(case.formatted_name)
        == "Иванов Иван Иванович - 20.10.2026.xlsx"
    )


@pytest.mark.parametrize("side", [DocumentSide.FRONT, DocumentSide.BACK, None])
def test_new_client_clears_missing_personal_fields_and_missing_photo(setup_copy, tmp_path, side):
    setup_copy = old_client(setup_copy, tmp_path)
    source, case, profile = setup_copy
    case.data.full_name = "НОВЫЙ КЛИЕНТ"
    if side:
        path = tmp_path / "new.png"
        Image.new("RGB", (600, 400), "blue").save(path)
        case.images.append(SourceImage(path=path, role=side))
    digest = file_hash(source)
    service = WorkbookService()
    plan = service.plan(source, case, profile)
    assert plan.identity_conflict
    with pytest.raises(AppError, match="смена клиента"):
        plan.select_missing_policy(MissingValuePolicy.PRESERVE_EXISTING)
    with pytest.raises(AppError, match="итоговой сводке"):
        service.write(plan, tmp_path / "blocked.xlsx", case, profile)
    plan.select_missing_policy(MissingValuePolicy.CLEAR_OLD_PERSONAL)
    plan.image_replace_confirmed = True
    output = service.write(plan, tmp_path / "new-client.xlsx", case, profile)
    assert file_hash(source) == digest
    book = load_workbook(output)
    assert book["Дог."]["A18"].value == "Новый Клиент"
    for cell in ["C22", "G19", "J20"]:
        assert book["Дог."][cell].value is None
    assert book["Дог."]["A21"].value == "=C20"
    assert len(book[profile.photo_sheet]._images) == (2 if side else 1)  # unrelated logo survives
    book.close()
    with ZipFile(source) as old, ZipFile(output) as new:
        for part in [
            "xl/worksheets/sheet2.xml",
            "xl/worksheets/sheet3.xml",
            "xl/worksheets/sheet4.xml",
            "xl/media/image3.png",
        ]:
            assert old.read(part) == new.read(part)


@pytest.mark.parametrize(
    "policy", [MissingValuePolicy.PRESERVE_EXISTING, MissingValuePolicy.CLEAR_OLD_PERSONAL]
)
def test_unknown_identity_requires_one_explicit_policy(setup_copy, tmp_path, policy):
    setup_copy = old_client(setup_copy, tmp_path)
    source, case, profile = setup_copy
    service = WorkbookService()
    plan = service.plan(source, case, profile)
    assert not plan.identity_conflict and plan.needs_missing_policy
    with pytest.raises(AppError, match="итоговой сводке"):
        service.write(plan, tmp_path / "blocked.xlsx", case, profile)
    plan.select_missing_policy(policy)
    output = service.write(plan, tmp_path / "chosen.xlsx", case, profile)
    book = load_workbook(output)
    clear = policy == MissingValuePolicy.CLEAR_OLD_PERSONAL
    assert book["Дог."]["C22"].value == (None if clear else "00123456789012")
    assert len(book[profile.photo_sheet]._images) == (1 if clear else 3)
    book.close()


def test_same_client_partial_photo_preserves_other_side_after_choice(setup_copy, tmp_path):
    setup_copy = old_client(setup_copy, tmp_path)
    source, case, profile = setup_copy
    case.data.full_name = "ПРЕЖНИЙ КЛИЕНТ"
    path = tmp_path / "front.png"
    Image.new("RGB", (800, 500), "blue").save(path)
    case.images = [SourceImage(path=path, role=DocumentSide.FRONT)]
    plan, output = save(setup_copy, tmp_path)
    assert not plan.identity_conflict
    assert len(plan.photos.insertions) == 1 and plan.photos.preserved_ids
    book = load_workbook(output)
    assert book["Дог."]["C22"].value == "00123456789012"
    assert len(book[profile.photo_sheet]._images) == 3
    book.close()


@pytest.mark.parametrize("status", [FieldStatus.INVALID, FieldStatus.CONFLICT])
def test_unusable_candidate_is_skipped_without_blocking(setup_copy, tmp_path, status):
    source, case, profile = setup_copy
    case.data.personal_identifier = "00123456789012"
    case.evidence["personal_identifier"] = FieldEvidence(status=status)
    assert not readiness(case, profile)
    plan, output = save(setup_copy, tmp_path)
    assert not any(c.field_name == "personal_identifier" for c in plan.changes)
    assert "personal_identifier" in plan.warnings.missing_fields


def test_formula_with_old_literal_cannot_be_cleared(setup_copy, tmp_path):
    source, case, profile = old_client(setup_copy, tmp_path)
    book = load_workbook(source)
    book["Дог."]["C22"] = '=CONCAT("00123456789012")'
    book.save(source)
    book.close()
    case.data.full_name = "НОВЫЙ КЛИЕНТ"
    with pytest.raises(AppError, match="формуле"):
        WorkbookService().plan(source, case, profile)


def test_no_formula_is_overwritten_with_new_ocr(setup_copy):
    source, case, profile = setup_copy
    book = load_workbook(source)
    book["Дог."]["A18"] = '="Прежний Клиент"'
    book.save(source)
    book.close()
    case.data.full_name = "НОВЫЙ КЛИЕНТ"
    with pytest.raises(AppError, match="формулу"):
        WorkbookService().plan(source, case, profile)


def test_copy_fallback_and_existing_result_never_overwritten(setup_copy, tmp_path, monkeypatch):
    def unsupported(*args):
        raise OSError(errno.EOPNOTSUPP, "no hard links")

    monkeypatch.setattr("document_excel.infrastructure.excel.os.link", unsupported)
    source, case, profile = setup_copy
    _, output = save(setup_copy, tmp_path)
    assert file_hash(output) == file_hash(source)
    with pytest.raises(OutputExistsError):
        save(setup_copy, tmp_path)
    assert not list(tmp_path.glob(".document-excel-*"))


def test_source_changed_and_mutated_plan_are_rejected(setup_copy, tmp_path):
    source, case, profile = setup_copy
    service = WorkbookService()
    case.data.full_name = "ИВАНОВ ИВАН"
    plan = service.plan(source, case, profile)
    plan.select_missing_policy(MissingValuePolicy.PRESERVE_EXISTING)
    plan.changes[0].cell = "N20"
    with pytest.raises(AppError, match="неожиданное"):
        service.write(plan, tmp_path / "tampered.xlsx", case, profile)
    plan = service.plan(source, case, profile)
    source.write_bytes(source.read_bytes() + b"changed")
    with pytest.raises(AppError, match="изменился"):
        service.write(plan, tmp_path / "changed.xlsx", case, profile)


def test_minimal_custom_mapping_does_not_require_pin_cells(tmp_path):
    source = tmp_path / "custom.xlsx"
    create_demo_template(source)
    service = WorkbookService()
    profile = legacy_profile().model_copy(
        update={"field_mappings": [CellMapping(field="full_name", sheet="Дог.", cell="A18")]}
    )
    profile = service.bind_profile(source, profile)
    case = ClientCase(
        workflow_mode=WorkflowMode.BEST_EFFORT_UPDATE,
        template_path=source,
        workbook_profile=profile.model_dump(mode="json"),
        data=ClientData(full_name="ИВАНОВ ИВАН"),
    )
    plan = service.plan(source, case, profile)
    assert len(plan.changes) == 1
    service.write(plan, tmp_path / "custom-copy.xlsx", case, profile)


def test_default_mode_and_partial_name_preserve_same_client(setup_copy, tmp_path):
    assert ClientCase().workflow_mode == WorkflowMode.BEST_EFFORT_UPDATE
    source, case, profile = old_client(setup_copy, tmp_path)
    case.data.full_name = "Прежний"
    plan, result = save(setup_copy, tmp_path)
    assert not plan.identity_conflict
    book = load_workbook(result)
    assert book["Дог."]["C22"].value == "00123456789012"
    book.close()


def test_stale_formula_cache_is_removed_without_changing_formula(setup_copy, tmp_path):
    from document_excel.infrastructure.xlsx_package import MAIN, child, document, encoded

    source, case, profile = old_client(setup_copy, tmp_path)
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    doc = document(parts["xl/worksheets/sheet1.xml"])
    for cell in doc.getElementsByTagNameNS(MAIN, "c"):
        if cell.getAttribute("r") == "I21":
            cell.setAttribute("t", "str")
            for node in list(cell.childNodes):
                if node.localName == "v":
                    cell.removeChild(node)
            value = child(cell, "v")
            value.appendChild(doc.createTextNode("Прежняя улица"))
            cell.appendChild(value)
    parts["xl/worksheets/sheet1.xml"] = encoded(doc)
    with ZipFile(source, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    case.data.full_name = "НОВЫЙ КЛИЕНТ"
    _, result = save(setup_copy, tmp_path, MissingValuePolicy.CLEAR_OLD_PERSONAL)
    book = load_workbook(result, data_only=True)
    assert book["Дог."]["I21"].value is None
    book.close()
    book = load_workbook(result)
    assert book["Дог."]["I21"].value == "=J20"
    book.close()


def test_empty_policy_cannot_silently_retain_old_client(setup_copy, tmp_path):
    source, case, profile = old_client(setup_copy, tmp_path)
    plan = WorkbookService().plan(source, case, profile)
    with pytest.raises(AppError, match="того же клиента"):
        plan.select_missing_policy(MissingValuePolicy.LEAVE_EMPTY)


def test_unknown_layout_is_blocked_even_without_ocr_values(tmp_path):
    from openpyxl import Workbook

    source = tmp_path / "unknown.xlsx"
    book = Workbook()
    book.active.title = "Дог."
    book.active["P30"] = "Unrelated layout"
    book.save(source)
    book.close()
    _, profile = WorkbookService().describe(source, legacy_profile(), strict=True)
    case = ClientCase(template_path=source)
    with pytest.raises(AppError, match="расположение полей"):
        WorkbookService().plan(source, case, profile)


def test_formula_cycle_cannot_retain_old_cache_for_new_client(setup_copy, tmp_path):
    source, case, profile = old_client(setup_copy, tmp_path)
    book = load_workbook(source)
    book["Дог."]["A21"] = "=I21"
    book["Дог."]["I21"] = "=A21"
    book.save(source)
    book.close()
    case.data.full_name = "НОВЫЙ КЛИЕНТ"
    with pytest.raises(AppError, match="формуле"):
        WorkbookService().plan(source, case, profile)


def test_fallback_write_failure_cleans_only_its_own_file(setup_copy, tmp_path, monkeypatch):
    source, case, profile = setup_copy
    case.data.full_name = "ИВАНОВ ИВАН"  # Use package writer, not the unchanged-copy shortcut.
    digest = file_hash(source)
    monkeypatch.setattr(
        "document_excel.infrastructure.excel.os.link",
        lambda *a: (_ for _ in ()).throw(OSError(errno.EOPNOTSUPP, "unsupported")),
    )

    def full_disk(source, target):
        target.write(b"incomplete")
        raise OSError(errno.ENOSPC, "full disk")

    monkeypatch.setattr("document_excel.infrastructure.excel.shutil.copyfileobj", full_disk)
    with pytest.raises(AppError, match="сохранить Excel"):
        save(setup_copy, tmp_path)
    assert file_hash(source) == digest
    assert not (tmp_path / "result.xlsx").exists()
    assert not list(tmp_path.glob(".document-excel-*"))


def test_fallback_publication_race_preserves_other_file(setup_copy, tmp_path, monkeypatch):
    def race(source, destination):
        destination.write_bytes(b"other complete file")
        raise OSError(errno.EOPNOTSUPP, "unsupported")

    monkeypatch.setattr("document_excel.infrastructure.excel.os.link", race)
    with pytest.raises(OutputExistsError):
        save(setup_copy, tmp_path)
    assert (tmp_path / "result.xlsx").read_bytes() == b"other complete file"
